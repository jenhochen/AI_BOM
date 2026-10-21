#!/usr/bin/env python3
"""
bom_matcher_0.06.py  v0.06
Text BOM (Altium) + Excel Parts DB  →  ADATA E-BOM formatted Excel

雙擊執行 / 無參數：啟動 GUI（自動偵測同目錄 TXT / XLSX / pcb.jpg|png / Knowledge XLSX）
帶參數執行：CLI 模式

Usage (CLI)
-----------
    python bom_matcher_0.06.py <text_bom.txt> <parts_db.xlsx> [-o out.xlsx]
"""

import re
import sys
import argparse
import threading
from pathlib import Path
from datetime import date

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext
    _HAS_TK = True
except ImportError:
    _HAS_TK = False

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import range_boundaries

VERSION = "0.06"

_KNOWN_PKGS = {
    '0201','0402','0603','0805','1206',
    'SOD','SOT-723','SOT723','DFN',
    'BGA-247','BGA247','BGA-154','BGA154',
}
# RC class: resistors (6103*) and capacitors (6104*) only
_RC_PREFIXES = ('6103', '6104')

# Description-based RC detection patterns (for items with no P/N)
_CAP_VAL_START = re.compile(r'^\d+\.?\d*\s*(pF|nF|uF)', re.IGNORECASE)
_RES_WORD      = re.compile(r'\bRES\b',                  re.IGNORECASE)
_OHM_VAL       = re.compile(r'\d+\.?\d*\s*(Ω|ohm)',      re.IGNORECASE)


# ════════════════════════════════════════════════════════════════════════════════
# 1. Parse Altium text BOM
# ════════════════════════════════════════════════════════════════════════════════

def parse_text_bom(path: str) -> list[dict]:
    items, cur = [], None
    with open(path, encoding='utf-8', errors='replace') as f:
        in_data = False
        for raw in f:
            line = raw.rstrip('\n')
            if 'Item' in line and 'Quantity' in line and 'Reference' in line:
                in_data = True; continue
            if not in_data or line.strip().startswith('__'):
                continue
            if line.startswith('\t\t') and cur:
                extra = line.strip().rstrip(',')
                if extra:
                    cur['ref'] = cur['ref'].rstrip(',') + ',' + extra
                continue
            cols = line.split('\t')
            if len(cols) >= 5 and cols[0].strip().isdigit():
                if cur: items.append(cur)
                cur = {
                    'item': int(cols[0].strip()),
                    'qty':  int(cols[1].strip()),
                    'ref':  cols[2].strip().rstrip(','),
                    'part': cols[3].strip(),
                    'fp':   cols[4].strip(),
                }
    if cur: items.append(cur)
    return items


# ════════════════════════════════════════════════════════════════════════════════
# 2. Read parts database
# ════════════════════════════════════════════════════════════════════════════════

def read_db(path: str, sheet=None) -> pd.DataFrame:
    kw = {'sheet_name': sheet} if sheet else {'sheet_name': 0}
    df = pd.read_excel(path, **kw)
    df['P/N']         = df['P/N'].astype(str).str.strip()
    df['Description'] = df['Description'].astype(str).str.strip()
    return df


# ════════════════════════════════════════════════════════════════════════════════
# 3. NC helpers
# ════════════════════════════════════════════════════════════════════════════════

def is_nc(part: str) -> bool:
    return any(s.strip().upper() == 'NC' for s in re.split(r'[/\\]+', part))

def strip_nc(part: str) -> str:
    segs = [s for s in re.split(r'[/\\]+', part) if s.strip().upper() != 'NC']
    return '/'.join(segs).strip('/')


# ════════════════════════════════════════════════════════════════════════════════
# 4. Value & package helpers
# ════════════════════════════════════════════════════════════════════════════════

_RES_PATTERNS = [
    (re.compile(r'([\d.]+)\s*R\b', re.I), r'\1Ω'),
    (re.compile(r'([\d.]+)\s*K\b', re.I), r'\1KΩ'),
    (re.compile(r'([\d.]+)\s*M\b', re.I), r'\1MΩ'),
]

def normalize(val: str) -> str:
    v = re.sub(r'_.*$', '', val).strip()
    for pat, rep in _RES_PATTERNS:
        v2, n = pat.subn(rep, v)
        if n: return v2
    return v

def pkg_of(fp: str) -> str:
    m = re.search(r'(0\d{3})', fp)
    if m: return m.group(1)
    fu = fp.upper()
    for tag in ('SOD','SOT-723','SOT723','DFN','BGA-247','BGA247','BGA-154','BGA154'):
        if tag in fu: return tag
    return ''

def pkg_ok(pkg: str, desc: str) -> bool:
    if not pkg: return True
    return pkg.replace('-','').upper() in desc.replace('-','').upper()

def value_in_desc(val: str, desc: str) -> bool:
    pattern = r'(?<![0-9A-Za-z.])' + re.escape(val) + r'(?![0-9A-Za-z])'
    return bool(re.search(pattern, desc, re.IGNORECASE))


# ════════════════════════════════════════════════════════════════════════════════
# 5. Core matching
# ════════════════════════════════════════════════════════════════════════════════

def find_matches(part: str, fp: str, db: pd.DataFrame) -> list:
    val  = strip_nc(part)
    pkg  = pkg_of(fp)
    vu   = val.upper()
    nrm  = normalize(val)
    fpt  = fp[0].upper() if fp else ''
    seen, matches = set(), []

    def add(row):
        k = str(row['P/N'])
        if k not in seen:
            seen.add(k); matches.append(row)

    for _, row in db.iterrows():
        desc = str(row['Description'])
        du   = desc.upper()

        if fpt == 'C':
            if not pkg_ok(pkg, desc): continue
            if value_in_desc(val, desc): add(row)

        elif fpt == 'R':
            if not pkg_ok(pkg, desc): continue
            if value_in_desc(val, desc) or value_in_desc(nrm, desc): add(row)

        elif fpt == 'L':
            if not pkg_ok(pkg, desc): continue
            if 'BEAD' in vu and 'BEAD' in du:
                m = re.search(r'(\d+)\s*[Oo]hm', val)
                if m and f"{m.group(1)}Ω" in desc: add(row)
            else:
                m = re.match(r'([\d.]+uH)', val, re.I)
                if m and m.group(1).upper() in du and 'INDUCTOR' in du: add(row)

        else:
            if vu == 'TVS':
                if 'TVS' in du:
                    if 'SOD' not in du and pkg in ('0201','0402'): add(row)
                    elif 'SOD' in du and 'SOD' in fp.upper(): add(row)
                continue
            if value_in_desc(val, desc):
                add(row); continue
            tokens = [t for t in re.split(r'[_\-/\s]+', val) if len(t) >= 4]
            if any(t.upper() in du for t in tokens): add(row)

    return matches


def match_item(item: dict, db: pd.DataFrame) -> tuple:
    """Returns (pn_str, desc_str, is_nc). Alternatives joined by ';'."""
    nc   = is_nc(item['part'])
    hits = find_matches(item['part'], item['fp'], db)

    if not hits:
        v = strip_nc(item['part'])
        p = pkg_of(item['fp'])
        display = p if p in _KNOWN_PKGS else item['fp']
        d = f"{v} {display}".strip()
        return '', d + (' (原本為NC)' if nc else ''), nc

    pn = str(hits[0]['P/N'])
    for h in hits[1:]:
        pn += f";{h['P/N']}"

    d = str(hits[0]['Description']).strip()
    return pn, d + (' (原本為NC)' if nc else ''), nc


def classify_rcec(pn: str, desc: str = '', fp: str = '') -> str:
    """
    RC: P/N starts with 6103 (resistor) or 6104 (capacitor).
    If no P/N, detect from description text or footprint prefix (C/R).
    All else → EC.
    """
    first = pn.split(';')[0].strip()
    if first.startswith(_RC_PREFIXES):
        return 'RC'
    if not first:
        if _CAP_VAL_START.search(desc): return 'RC'   # e.g. "10pF 0201"
        if _RES_WORD.search(desc):      return 'RC'   # e.g. "RES 22KΩ 0201"
        if _OHM_VAL.search(desc):       return 'RC'   # e.g. "300Ω 0201"
        if fp and fp[0].upper() in ('C', 'R'): return 'RC'  # C0201 / R0201
    return 'EC'


# ════════════════════════════════════════════════════════════════════════════════
# 6. Knowledge loading & matching
# ════════════════════════════════════════════════════════════════════════════════

def load_knowledge(path: str) -> list[dict]:
    """Load Layout_and_BOM_Knowledge file → list of {pn, 品名, 包裝, note}."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    entries = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 5:
            continue
        _, pn, name, pkg, note = row[0], row[1], row[2], row[3], row[4]
        if not any([pn, name, pkg]):
            continue
        if isinstance(pn, (int, float)):
            pn = str(int(pn))
        entries.append({
            'pn':   str(pn).strip()   if pn   else '',
            '品名': str(name).strip() if name else '',
            '包裝': str(pkg).strip()  if pkg  else '',
            'note': str(note).strip() if note else '',
        })
    return entries


def find_knowledge_hits(pn_str: str, desc_str: str, knowledge: list) -> str:
    """
    Return merged note (newline-separated, deduplicated) for all matching
    entries. Match triggers if any ONE of: PN / 品名 / 包裝 matches.
    """
    pn_str   = str(pn_str  or '')
    desc_str = str(desc_str or '')
    bom_pns  = {p.strip() for p in pn_str.split(';') if p.strip()}
    desc_nd  = desc_str.replace('-', '').upper()

    seen_notes, notes = set(), []
    for k in knowledge:
        hit = False
        if k['pn'] and k['pn'] in bom_pns:
            hit = True
        if not hit and k['品名'] and k['品名'].upper() in desc_str.upper():
            hit = True
        if not hit and k['包裝']:
            if k['包裝'].replace('-', '').upper() in desc_nd:
                hit = True
        if hit and k['note'] and k['note'] not in seen_notes:
            seen_notes.add(k['note'])
            notes.append(k['note'])
    return '\n'.join(notes)


# ════════════════════════════════════════════════════════════════════════════════
# 7. Style helpers
# ════════════════════════════════════════════════════════════════════════════════

def _sd(style='thin'):
    return Side(style=style, color='000000')

def _brd(l='thin', r='thin', t='thin', b='thin') -> Border:
    def s(x): return _sd(x) if x else Side()
    return Border(left=s(l), right=s(r), top=s(t), bottom=s(b))

def _aln(h='left', v='center', wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


B_OUTER  = _brd('medium', 'medium', 'thin', 'thin')
B_TOPCAP = _brd('medium', 'medium', 'medium', 'thin')
B_BOTCAP = _brd('medium', 'medium', 'thin', 'medium')
B_BODY_L = _brd('medium', 'medium', 'thin', 'thin')
B_FOOT   = _brd('medium', 'medium', 'thin', 'thin')


def _set_box(ws, rng: str, value=None, font=None, align=None,
             fill=None, border: Border = None):
    """
    Merge cells (if range), distribute border correctly across all edge
    cells, and write value/font/align/fill to the anchor cell.
    Excel renders merged-cell borders from individual edge cells, not
    just the anchor — this function ensures every boundary is drawn.
    """
    c1, r1, c2, r2 = range_boundaries(rng)
    if (c1, r1) != (c2, r2):
        ws.merge_cells(rng)

    if border is not None:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = Border(
                    left=border.left   if c == c1 else Side(),
                    right=border.right  if c == c2 else Side(),
                    top=border.top    if r == r1 else Side(),
                    bottom=border.bottom if r == r2 else Side(),
                )
    if fill is not None:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = fill

    anchor = ws.cell(r1, c1)
    if value is not None: anchor.value     = value
    if font  is not None: anchor.font      = font
    if align is not None: anchor.alignment = align
    return anchor


# ════════════════════════════════════════════════════════════════════════════════
# 8. Header (rows 1-8)
# ════════════════════════════════════════════════════════════════════════════════

def _build_header(ws):
    FA = lambda sz, bold=False: Font(name='Arial',  size=sz, bold=bold)
    FK = lambda sz:             Font(name='標楷體', size=sz)
    R, L, C = _aln('right'), _aln('left'), _aln('center')

    ws.merge_cells('A1:O1')
    _set_box(ws, 'A1', 'ADATA Technology Co., Ltd.', FA(14, True), L)
    ws.row_dimensions[1].height = 17.4

    ws.merge_cells('A2:L2')
    _set_box(ws, 'A2', 'M.2 SSD E-BOM', FA(24, True), C)
    ws.row_dimensions[2].height = 30.6

    ws.row_dimensions[3].height = 19.8
    _set_box(ws, 'A3:C3', 'PCB板號：',  FA(14), R, border=B_TOPCAP)
    _set_box(ws, 'D3:J3', None,         FA(14), L, border=B_TOPCAP)
    _set_box(ws, 'K3:L3', '發行序號：', FK(14), R, border=B_TOPCAP)
    _set_box(ws, 'M3:O3', None,         FA(14), L, border=B_TOPCAP)

    std = [
        (4, 19.8, 'PCB料號：',   '發行(改版)日期：', None),
        (5, 39.6, 'RC/EC NO.：', '目前版次：',        '0'),
        (6, 19.8, 'PCB供應商：', None,                None),
    ]
    for row, h, lbl_ac, lbl_kl, val_mo in std:
        ws.row_dimensions[row].height = h
        _set_box(ws, f'A{row}:C{row}', lbl_ac, FA(14), R, border=B_OUTER)
        _set_box(ws, f'D{row}:J{row}', None,   FA(14), L, border=B_OUTER)
        _set_box(ws, f'K{row}:L{row}', lbl_kl, FK(14), R, border=B_OUTER)
        _set_box(ws, f'M{row}:O{row}', val_mo, FA(14), L, border=B_OUTER)

    ws.row_dimensions[7].height = 19.8
    ws.row_dimensions[8].height = 20.4
    _set_box(ws, 'A7:C7', 'PCB 規格：',       FA(14), R, border=B_OUTER)
    _set_box(ws, 'D7:J7', None,               FA(14), L, border=B_OUTER)
    _set_box(ws, 'A8:C8', 'Support IC Type：', FA(14), R, border=B_BOTCAP)
    _set_box(ws, 'D8:J8', None,               FA(14), L, border=B_BOTCAP)

    box78 = _brd('medium', 'medium', 'thin', 'medium')
    _set_box(ws, 'K7:L8', '環保等級：', FK(14), R, border=box78)
    _set_box(ws, 'M7:O8',
            '□A級 □S級 ■H級 □G級\n□OTHER 特殊客戶要求  :  ',
            FK(14), _aln('left', 'center', True), border=box78)


# ════════════════════════════════════════════════════════════════════════════════
# 9. PCB image (rows 9-35)
# ════════════════════════════════════════════════════════════════════════════════

def _add_pcb_image(ws, img_path):
    for r in range(9, 36):
        ws.row_dimensions[r].height = 28.2
        _set_box(ws, f'A{r}', None, border=_brd('medium', None, None, None))
        _set_box(ws, f'O{r}', None, border=_brd(None, 'medium', None, None))

    if not img_path:
        return
    p = Path(img_path)
    if not p.exists():
        print(f"[警告] 找不到 PCB 圖片: {img_path}")
        return
    try:
        img = XLImage(str(p))
        orig_w, orig_h = img.width, img.height
        target_h = 700
        scale = target_h / orig_h
        img.height = target_h
        img.width  = int(orig_w * scale)
        ws.add_image(img, 'A9')
        print(f"✓ PCB 圖片已嵌入: {p.name}  ({img.width}×{img.height}px)")
    except Exception as e:
        print(f"[警告] PCB 圖片嵌入失敗: {e}")


# ════════════════════════════════════════════════════════════════════════════════
# 10. BOM table
# ════════════════════════════════════════════════════════════════════════════════

def _build_bom_header(ws, row: int):
    """Row 36 column headers. G:O = 'Knowledge Note' (single-side PCB)."""
    F  = Font(name='Arial', size=12, bold=True)
    BC = _aln('center')
    B  = _brd('medium', 'medium', 'medium', 'thin')
    for col, val in [('A','Item'),('B','RC/EC'),('C','P/N'),
                     ('D','Description'),('E','Qty'),
                     ('F','Location-Top Layer')]:
        _set_box(ws, f'{col}{row}', val, F, BC, border=B)
    _set_box(ws, f'G{row}:O{row}', 'Knowledge Note', F, BC, border=B)
    ws.row_dimensions[row].height = 15.0


FL_NC  = PatternFill('solid', start_color='FFF2CC')   # yellow  = NC
FL_ALT = PatternFill('solid', start_color='E2EFDA')   # green   = alt P/N
FL_KN  = PatternFill('solid', start_color='FFE0B2')   # orange  = knowledge hit
FL_TOT = PatternFill('solid', start_color='D9E1F2')   # blue    = total


def _write_data_row(ws, row: int, seq: int, rcec: str,
                    pn: str, desc: str, qty, ref: str,
                    nc: bool, kn_note: str = ''):
    F_A = Font(name='Arial', size=14)
    F_B = Font(name='Arial', size=12)
    F_N = Font(name='Arial', size=10)

    # Fill priority: knowledge orange > NC yellow > alt green
    fill = (FL_KN if kn_note
            else FL_NC if nc
            else FL_ALT if ';' in pn
            else None)

    data = [
        ('A', seq,  F_A, _aln('center')),
        ('B', rcec, F_B, _aln('center')),
        ('C', pn,   F_B, _aln('left')),
        ('D', desc, F_B, _aln('left', wrap=True)),
        ('E', qty,  F_B, _aln('center')),
        ('F', ref,  F_B, _aln('left', wrap=True)),
    ]
    for col, val, fnt, aln in data:
        _set_box(ws, f'{col}{row}', val, fnt, aln, fill=fill, border=B_BODY_L)

    # G:O = Knowledge Note
    note_val = kn_note if kn_note else None
    _set_box(ws, f'G{row}:O{row}', note_val, F_N,
             _aln('left', wrap=True), fill=fill, border=B_BODY_L)

    lines = max(1, kn_note.count('\n') + 1) if kn_note else 1
    ws.row_dimensions[row].height = max(17.4, lines * 18.0)


def _write_total_row(ws, row: int, total_qty: int):
    F = Font(name='Arial', size=12, bold=True)
    _set_box(ws, f'A{row}:B{row}', 'Total', F, _aln('center'),
             fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'C{row}', None, F, _aln('center'), fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'D{row}', None, F, _aln('center'), fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'E{row}', total_qty, F, _aln('center'),
             fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'F{row}:O{row}', None, F, _aln('center'),
             fill=FL_TOT, border=B_BODY_L)
    ws.row_dimensions[row].height = 17.4


def _write_footer(ws, row: int):
    F_SM  = Font(name='Arial', size=10)
    F_LBL = Font(name='Arial', size=11)

    row += 1
    _set_box(ws, f'A{row}:O{row}', None, border=B_FOOT)

    row += 1
    _set_box(ws, f'A{row}:O{row}',
             'Remark：Top Layer only BOM. NC items listed for reference.',
             F_SM, _aln('left'), border=B_FOOT)
    ws.row_dimensions[row].height = 14.0

    row += 1
    _set_box(ws, f'A{row}:O{row}', None, border=B_FOOT)

    row += 1
    _set_box(ws, f'A{row}', '圖例',
             Font(name='Arial', bold=True, size=9), border=B_FOOT)
    _set_box(ws, f'B{row}:O{row}', None, border=B_FOOT)

    legends = [
        ('黃底', FL_NC,  '原本為 NC 的零件'),
        ('綠底', FL_ALT, '含替代料 (P/N 以「;」分隔)'),
        ('橘底', FL_KN,  'Knowledge 命中（含注意事項）'),
    ]
    for i, (label, fl, note) in enumerate(legends):
        r = row + 1 + i
        _set_box(ws, f'A{r}', label,
                 Font(name='Arial', size=9), fill=fl, border=B_FOOT)
        _set_box(ws, f'B{r}', note,
                 Font(name='Arial', size=9), border=B_FOOT)
        _set_box(ws, f'C{r}:O{r}', None, border=B_FOOT)
    row += len(legends)

    row += 1
    _set_box(ws, f'A{row}:O{row}', None, border=B_FOOT)   # blank divider
    row += 1
    _set_box(ws, f'A{row}:B{row+1}', '    Approval  :  ', F_LBL,
             _aln('center'), border=B_FOOT)
    _set_box(ws, f'C{row}:D{row+1}', None, border=B_FOOT)
    _set_box(ws, f'E{row}:F{row+1}', 'Rechecked: ', F_LBL,
             _aln('center'), border=B_FOOT)
    _set_box(ws, f'G{row}:H{row+1}', 'Prepared: ', F_LBL,
             _aln('center'), border=B_FOOT)
    _set_box(ws, f'I{row}:J{row+1}', None, border=B_FOOT)
    _set_box(ws, f'K{row}:O{row+1}', None, border=B_FOOT)

    last = row + 1
    for col in range(1, 16):
        c = ws.cell(last, col)
        c.border = Border(left=c.border.left, right=c.border.right,
                          top=c.border.top, bottom=_sd('medium'))


def _set_col_widths(ws):
    for col, w in {
        'A': 5.1,  'B': 10.1, 'C': 29.1, 'D': 77.8, 'E': 9.2,
        'F': 54.3, 'G': 11.0, 'H': 16.0, 'I': 13.0, 'J': 13.0,
        'K':  9.0, 'L': 12.0, 'M': 13.3, 'N':  9.0, 'O': 13.0,
    }.items():
        ws.column_dimensions[col].width = w


# ════════════════════════════════════════════════════════════════════════════════
# 11. Main write function
# ════════════════════════════════════════════════════════════════════════════════

def write_excel(items: list, db: pd.DataFrame,
                out_path: str, pcb_img: str = None,
                knowledge: list = None):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'BOM'
    ws.sheet_view.showGridLines = False

    _set_col_widths(ws)
    _build_header(ws)
    _add_pcb_image(ws, pcb_img)

    BOM_HDR = 36
    _build_bom_header(ws, BOM_HDR)

    total_qty = 0
    data_row  = BOM_HDR + 1

    for seq, item in enumerate(items, 1):
        pn, desc, nc = match_item(item, db)
        qty, ref = item['qty'], item['ref']
        total_qty += qty
        rcec    = classify_rcec(pn, desc, item['fp'])
        kn_note = find_knowledge_hits(pn, desc, knowledge) if knowledge else ''
        _write_data_row(ws, data_row, seq, rcec, pn, desc, qty, ref, nc, kn_note)
        data_row += 1

    # ── PCB row: last item, EC, Description fixed = 'PCB', Qty = 1 ──────────
    pcb_seq  = len(items) + 1
    kn_pcb   = find_knowledge_hits('', 'PCB', knowledge) if knowledge else ''
    _write_data_row(ws, data_row, pcb_seq, 'EC', '', 'PCB', 1, '', False, kn_pcb)
    total_qty += 1
    data_row  += 1

    _write_total_row(ws, data_row, total_qty)
    _write_footer(ws, data_row)

    # ── v0.06 fix: no freeze panes (PCB image was frozen in v0.05) ───────────
    ws.freeze_panes = None

    wb.save(out_path)
    print(f"✓ 輸出: {out_path}  |  {len(items)+1} 項(含PCB)  |  Total Qty = {total_qty}")


# ════════════════════════════════════════════════════════════════════════════════
# 12. Auto-detect helpers
# ════════════════════════════════════════════════════════════════════════════════

def _find_pcb_img(directory: str) -> str:
    d = Path(directory)
    for name in ('pcb.jpg','pcb.jpeg','pcb.png',
                 'PCB.jpg','PCB.JPG','PCB.png','PCB.PNG'):
        p = d / name
        if p.exists(): return str(p)
    return ''


def _find_knowledge(directory: str) -> str:
    """Auto-detect Layout*Knowledge*.xlsx or *Knowledge*.xlsx in directory."""
    d = Path(directory)
    for pattern in ('Layout*Knowledge*.xlsx',
                    '*Knowledge*.xlsx', '*knowledge*.xlsx'):
        matches = sorted(d.glob(pattern))
        if matches:
            return str(matches[-1])
    return ''


# ════════════════════════════════════════════════════════════════════════════════
# 13. GUI
# ════════════════════════════════════════════════════════════════════════════════

class _StdoutToText:
    def __init__(self, widget, root):
        self.widget = widget; self.root = root
    def write(self, s: str):
        def _ins():
            self.widget.configure(state='normal')
            self.widget.insert('end', s)
            self.widget.see('end')
            self.widget.configure(state='disabled')
        self.root.after(0, _ins)
    def flush(self): pass


class BomMatcherApp(tk.Tk if _HAS_TK else object):
    PAD = 8
    def __init__(self):
        super().__init__()
        self.title(f'BOM Matcher  v{VERSION}')
        self.resizable(True, True)
        self.minsize(740, 580)
        self._build_ui()
        self.after(100, self._auto_detect)

    def _build_ui(self):
        P = self.PAD
        self.configure(padx=P, pady=P)

        cwd_frame = tk.LabelFrame(self, text='目前工作目錄', padx=P, pady=4)
        cwd_frame.pack(fill='x', pady=(0, P))
        tk.Label(cwd_frame, text=str(Path.cwd()), anchor='w', fg='#333').pack(fill='x')

        file_frame = tk.LabelFrame(self, text='檔案設定', padx=P, pady=P)
        file_frame.pack(fill='x', pady=(0, P))
        file_frame.columnconfigure(1, weight=1)

        self.var_bom = tk.StringVar()
        self.var_db  = tk.StringVar()
        self.var_pcb = tk.StringVar()
        self.var_kn  = tk.StringVar()
        self.var_out = tk.StringVar()

        rows_cfg = [
            ('Text BOM (.txt)',        self.var_bom, self._browse_bom),
            ('Excel DB (.xlsx)',       self.var_db,  self._browse_db),
            ('PCB 圖片 (pcb.jpg)',     self.var_pcb, self._browse_pcb),
            ('Knowledge DB (.xlsx)',   self.var_kn,  self._browse_kn),
            ('輸出 BOM (.xlsx)',       self.var_out, self._browse_out),
        ]
        for r, (label, var, cmd) in enumerate(rows_cfg):
            tk.Label(file_frame, text=label, anchor='e', width=20)\
                .grid(row=r, column=0, sticky='e', pady=3, padx=(0,6))
            tk.Entry(file_frame, textvariable=var, width=56)\
                .grid(row=r, column=1, sticky='ew', pady=3)
            tk.Button(file_frame, text='瀏覽…', width=7, command=cmd)\
                .grid(row=r, column=2, padx=(6,0), pady=3)

        btn_frame = tk.Frame(self)
        btn_frame.pack(fill='x', pady=(0, P))
        self.btn_run = tk.Button(
            btn_frame, text='▶  執行', width=16,
            bg='#1F4E79', fg='white', font=('Arial',11,'bold'),
            relief='flat', cursor='hand2', command=self._run)
        self.btn_run.pack(side='right')

        log_frame = tk.LabelFrame(self, text='執行日誌', padx=P, pady=P)
        log_frame.pack(fill='both', expand=True)
        self.log_text = scrolledtext.ScrolledText(
            log_frame, state='disabled', height=12,
            font=('Consolas',9), bg='#1e1e1e', fg='#d4d4d4',
            insertbackground='white', wrap='word')
        self.log_text.pack(fill='both', expand=True)

    def _auto_detect(self):
        cwd   = Path.cwd()
        txts  = sorted(cwd.glob('*.txt'))
        xlsxs = sorted(cwd.glob('*.xlsx'))
        self._log(f'[掃描] 目錄：{cwd}\n')

        if len(txts) == 1:
            self._set_bom(str(txts[0]))
            self._log(f'[自動偵測] Text BOM    → {txts[0].name}')
        elif len(txts) > 1:
            self._log(f'[提示] 找到 {len(txts)} 個 TXT，請手動選擇：')
            for f in txts: self._log(f'          {f.name}')
        else:
            self._log('[提示] 未找到 TXT BOM，請手動選擇')

        # Separate knowledge xlsx from db xlsx
        kn_files  = sorted(cwd.glob('*Knowledge*.xlsx')) + \
                    sorted(cwd.glob('*knowledge*.xlsx'))
        db_files  = [f for f in xlsxs if f not in kn_files]

        if len(db_files) == 1:
            self.var_db.set(str(db_files[0]))
            self._log(f'[自動偵測] Excel DB    → {db_files[0].name}')
        elif len(db_files) > 1:
            self._log(f'[提示] 找到 {len(db_files)} 個 DB XLSX，請手動選擇：')
            for f in db_files: self._log(f'          {f.name}')
        else:
            self._log('[提示] 未找到 Excel DB，請手動選擇')

        if kn_files:
            self.var_kn.set(str(kn_files[-1]))
            self._log(f'[自動偵測] Knowledge   → {kn_files[-1].name}')
        else:
            self._log('[提示] 未找到 Knowledge DB（選填）')

        pcb = _find_pcb_img(str(cwd))
        if pcb:
            self.var_pcb.set(pcb)
            self._log(f'[自動偵測] PCB 圖片    → {Path(pcb).name}')
        else:
            self._log('[提示] 未找到 pcb.jpg/png（選填）')
        self._log('')

    def _set_bom(self, path: str):
        self.var_bom.set(path)
        if not self.var_out.get():
            p = Path(path)
            self.var_out.set(str(p.parent / f'BOM_{p.stem}_{date.today():%Y%m%d}.xlsx'))

    def _browse_bom(self):
        f = filedialog.askopenfilename(title='選擇 Text BOM', initialdir=Path.cwd(),
            filetypes=[('Text BOM','*.txt'),('All','*.*')])
        if f: self._set_bom(f)

    def _browse_db(self):
        f = filedialog.askopenfilename(title='選擇零件資料庫', initialdir=Path.cwd(),
            filetypes=[('Excel','*.xlsx'),('All','*.*')])
        if f: self.var_db.set(f)

    def _browse_pcb(self):
        f = filedialog.askopenfilename(title='選擇 PCB 圖片', initialdir=Path.cwd(),
            filetypes=[('Image','*.jpg *.jpeg *.png'),('All','*.*')])
        if f: self.var_pcb.set(f)

    def _browse_kn(self):
        f = filedialog.askopenfilename(title='選擇 Knowledge DB', initialdir=Path.cwd(),
            filetypes=[('Excel','*.xlsx'),('All','*.*')])
        if f: self.var_kn.set(f)

    def _browse_out(self):
        f = filedialog.asksaveasfilename(title='輸出 BOM 路徑', initialdir=Path.cwd(),
            defaultextension='.xlsx', filetypes=[('Excel','*.xlsx')])
        if f: self.var_out.set(f)

    def _run(self):
        bom = self.var_bom.get().strip()
        db  = self.var_db.get().strip()
        out = self.var_out.get().strip()
        if not bom: messagebox.showerror('錯誤','請選擇 Text BOM 檔案'); return
        if not db:  messagebox.showerror('錯誤','請選擇 Excel 零件資料庫'); return
        if not out: messagebox.showerror('錯誤','請指定輸出路徑'); return
        self.btn_run.configure(state='disabled')
        self._log('=' * 55)
        threading.Thread(
            target=self._run_task,
            args=(bom, db, self.var_pcb.get().strip(),
                  self.var_kn.get().strip(), out),
            daemon=True).start()

    def _run_task(self, bom, db_path, pcb, kn_path, out):
        old = sys.stdout
        sys.stdout = _StdoutToText(self.log_text, self)
        try:
            print(f'BOM Matcher v{VERSION}')
            print(f'  Text BOM    : {bom}')
            print(f'  Excel DB    : {db_path}')
            print(f'  PCB 圖片    : {pcb or "（未指定）"}')
            print(f'  Knowledge   : {kn_path or "（未指定）"}')
            print(f'  Output      : {out}\n')
            items = parse_text_bom(bom)
            print(f'✓ 解析 Text BOM: {len(items)} 項')
            db_df = read_db(db_path)
            print(f'✓ 讀取零件資料庫: {len(db_df)} 筆')
            kn = None
            if kn_path and Path(kn_path).exists():
                kn = load_knowledge(kn_path)
                print(f'✓ 載入 Knowledge: {len(kn)} 筆')
            write_excel(items, db_df, out, pcb or None, kn)
            self.after(0, lambda: messagebox.showinfo('完成', f'BOM 已產出：\n{out}'))
        except Exception as e:
            import traceback; traceback.print_exc()
            print(f'\n✗ 錯誤：{e}')
            self.after(0, lambda: messagebox.showerror('執行錯誤', str(e)))
        finally:
            sys.stdout = old
            self.after(0, lambda: self.btn_run.configure(state='normal'))
            self.after(0, lambda: self._log(''))

    def _log(self, msg: str):
        self.log_text.configure(state='normal')
        self.log_text.insert('end', msg + '\n')
        self.log_text.see('end')
        self.log_text.configure(state='disabled')


# ════════════════════════════════════════════════════════════════════════════════
# 14. Entry point
# ════════════════════════════════════════════════════════════════════════════════

def _cli():
    ap = argparse.ArgumentParser(description=f'BOM Matcher v{VERSION}')
    ap.add_argument('text_bom')
    ap.add_argument('excel_db')
    ap.add_argument('-o','--output')
    ap.add_argument('-s','--sheet')
    ap.add_argument('-k','--knowledge', help='Knowledge DB xlsx path')
    args = ap.parse_args()

    out = args.output or str(
        Path(args.text_bom).parent /
        f'BOM_{Path(args.text_bom).stem}_{date.today():%Y%m%d}.xlsx')
    pcb = _find_pcb_img(str(Path(args.text_bom).parent))
    kn_path = args.knowledge or _find_knowledge(str(Path(args.text_bom).parent))

    print(f'BOM Matcher v{VERSION}')
    print(f'  Text BOM  : {args.text_bom}')
    print(f'  Excel DB  : {args.excel_db}')
    print(f'  PCB 圖片  : {pcb or "（未找到）"}')
    print(f'  Knowledge : {kn_path or "（未找到）"}')
    print(f'  Output    : {out}\n')

    items = parse_text_bom(args.text_bom)
    print(f'✓ 解析 Text BOM: {len(items)} 項')
    db = read_db(args.excel_db, args.sheet)
    print(f'✓ 讀取零件資料庫: {len(db)} 筆')
    kn = None
    if kn_path and Path(kn_path).exists():
        kn = load_knowledge(kn_path)
        print(f'✓ 載入 Knowledge: {len(kn)} 筆')
    write_excel(items, db, out, pcb or None, kn)


def main():
    if len(sys.argv) > 1:
        _cli()
    else:
        if not _HAS_TK:
            print("錯誤：GUI 需要 tkinter，請改用 CLI 模式"); sys.exit(1)
        BomMatcherApp().mainloop()


if __name__ == '__main__':
    main()

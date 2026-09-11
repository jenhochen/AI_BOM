#!/usr/bin/env python3
"""
bom_matcher.py  v1.1
Text BOM (Altium) + Excel Parts DB  →  Formatted BOM Excel

雙擊執行 / 無參數：啟動 GUI 介面（自動偵測同目錄下的 TXT / XLSX）
帶參數執行：CLI 模式（同 v1.0）

Usage (CLI)
-----------
    python bom_matcher.py <text_bom.txt> <parts_db.xlsx> [-o out.xlsx] [-s SheetName]
"""

import re
import sys
import argparse
import threading
from pathlib import Path
from datetime import date

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext
    _HAS_TK = True
except ImportError:
    _HAS_TK = False

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

VERSION = "1.1"

_KNOWN_PKGS = {
    '0201','0402','0603','0805','1206',
    'SOD','SOT-723','SOT723','DFN',
    'BGA-247','BGA247','BGA-154','BGA154',
}


# ════════════════════════════════════════════════════════════════════════════════
# 1. Parse Altium text BOM
# ════════════════════════════════════════════════════════════════════════════════

def parse_text_bom(path: str) -> list[dict]:
    items, cur = [], None
    with open(path, encoding='utf-8', errors='replace') as f:
        in_data = False
        for raw in f:
            line = raw.rstrip('\n')
            if 'Item' in line and 'Quantity' in line and 'Reference' in line:
                in_data = True; continue
            if not in_data or line.strip().startswith('__'):
                continue
            if line.startswith('\t\t') and cur:
                extra = line.strip().rstrip(',')
                if extra:
                    cur['ref'] = cur['ref'].rstrip(',') + ',' + extra
                continue
            cols = line.split('\t')
            if len(cols) >= 5 and cols[0].strip().isdigit():
                if cur: items.append(cur)
                cur = {
                    'item': int(cols[0].strip()),
                    'qty':  int(cols[1].strip()),
                    'ref':  cols[2].strip().rstrip(','),
                    'part': cols[3].strip(),
                    'fp':   cols[4].strip(),
                }
    if cur: items.append(cur)
    return items


# ════════════════════════════════════════════════════════════════════════════════
# 2. Read parts database
# ════════════════════════════════════════════════════════════════════════════════

def read_db(path: str, sheet=None) -> pd.DataFrame:
    kw = {'sheet_name': sheet} if sheet else {'sheet_name': 0}
    df = pd.read_excel(path, **kw)
    df['P/N']         = df['P/N'].astype(str).str.strip()
    df['Description'] = df['Description'].astype(str).str.strip()
    return df


# ════════════════════════════════════════════════════════════════════════════════
# 3. NC helpers
# ════════════════════════════════════════════════════════════════════════════════

def is_nc(part: str) -> bool:
    return any(s.strip().upper() == 'NC' for s in re.split(r'[/\\]+', part))

def strip_nc(part: str) -> str:
    segs = [s for s in re.split(r'[/\\]+', part) if s.strip().upper() != 'NC']
    return '/'.join(segs).strip('/')


# ════════════════════════════════════════════════════════════════════════════════
# 4. Value & package helpers
# ════════════════════════════════════════════════════════════════════════════════

_RES_PATTERNS = [
    (re.compile(r'([\d.]+)\s*R\b', re.I), r'\1Ω'),
    (re.compile(r'([\d.]+)\s*K\b', re.I), r'\1KΩ'),
    (re.compile(r'([\d.]+)\s*M\b', re.I), r'\1MΩ'),
]

def normalize(val: str) -> str:
    v = re.sub(r'_.*$', '', val).strip()
    for pat, rep in _RES_PATTERNS:
        v2, n = pat.subn(rep, v)
        if n: return v2
    return v

def pkg_of(fp: str) -> str:
    m = re.search(r'(0\d{3})', fp)
    if m: return m.group(1)
    fu = fp.upper()
    for tag in ('SOD','SOT-723','SOT723','DFN','BGA-247','BGA247','BGA-154','BGA154'):
        if tag in fu: return tag
    return ''

def pkg_ok(pkg: str, desc: str) -> bool:
    if not pkg: return True
    return pkg.replace('-','').upper() in desc.replace('-','').upper()

def value_in_desc(val: str, desc: str) -> bool:
    pattern = r'(?<![0-9A-Za-z.])' + re.escape(val) + r'(?![0-9A-Za-z])'
    return bool(re.search(pattern, desc, re.IGNORECASE))


# ════════════════════════════════════════════════════════════════════════════════
# 5. Core matching
# ════════════════════════════════════════════════════════════════════════════════

def find_matches(part: str, fp: str, db: pd.DataFrame) -> list:
    val  = strip_nc(part)
    pkg  = pkg_of(fp)
    vu   = val.upper()
    nrm  = normalize(val)
    fpt  = fp[0].upper() if fp else ''
    seen, matches = set(), []

    def add(row):
        k = str(row['P/N'])
        if k not in seen:
            seen.add(k); matches.append(row)

    for _, row in db.iterrows():
        desc = str(row['Description'])
        du   = desc.upper()

        if fpt == 'C':
            if not pkg_ok(pkg, desc): continue
            if value_in_desc(val, desc): add(row)

        elif fpt == 'R':
            if not pkg_ok(pkg, desc): continue
            if value_in_desc(val, desc) or value_in_desc(nrm, desc): add(row)

        elif fpt == 'L':
            if not pkg_ok(pkg, desc): continue
            if 'BEAD' in vu and 'BEAD' in du:
                m = re.search(r'(\d+)\s*[Oo]hm', val)
                if m and f"{m.group(1)}Ω" in desc: add(row)
            else:
                m = re.match(r'([\d.]+uH)', val, re.I)
                if m and m.group(1).upper() in du and 'INDUCTOR' in du: add(row)

        else:
            if vu == 'TVS':
                if 'TVS' in du:
                    if 'SOD' not in du and pkg in ('0201','0402'): add(row)
                    elif 'SOD' in du and 'SOD' in fp.upper(): add(row)
                continue
            if value_in_desc(val, desc):
                add(row); continue
            tokens = [t for t in re.split(r'[_\-/\s]+', val) if len(t) >= 4]
            if any(t.upper() in du for t in tokens): add(row)

    return matches


def match_item(item: dict, db: pd.DataFrame) -> tuple:
    nc   = is_nc(item['part'])
    hits = find_matches(item['part'], item['fp'], db)

    if not hits:
        v = strip_nc(item['part'])
        p = pkg_of(item['fp'])
        display = p if p in _KNOWN_PKGS else item['fp']
        d = f"{v} {display}".strip()
        return '', d + (' (原本為NC)' if nc else ''), nc

    pn = str(hits[0]['P/N'])
    for h in hits[1:]: pn += f" (或 {h['P/N']})"
    d = str(hits[0]['Description']).strip()
    return pn, d + (' (原本為NC)' if nc else ''), nc


# ════════════════════════════════════════════════════════════════════════════════
# 6. Excel output
# ════════════════════════════════════════════════════════════════════════════════

def _border() -> Border:
    s = Side(style='thin', color='AAAAAA')
    return Border(left=s, right=s, top=s, bottom=s)

def write_excel(items: list, db: pd.DataFrame, out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'BOM'

    F_HDR  = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    FL_HDR = PatternFill('solid', start_color='1F4E79')
    F_BODY = Font(name='Arial', size=10)
    F_TOT  = Font(name='Arial', bold=True, size=10)
    FL_NC  = PatternFill('solid', start_color='FFF2CC')
    FL_ALT = PatternFill('solid', start_color='E2EFDA')
    FL_TOT = PatternFill('solid', start_color='D9E1F2')
    A_CTR  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    A_LFT  = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    BRD    = _border()

    for col, h in enumerate(['P/N','Description','Qty','Location-Top Layer'], 1):
        cell = ws.cell(1, col, h)
        cell.font = F_HDR; cell.fill = FL_HDR
        cell.alignment = A_CTR; cell.border = BRD
    ws.row_dimensions[1].height = 22

    total_qty = 0
    for r, item in enumerate(items, 2):
        pn, desc, nc = match_item(item, db)
        qty, ref = item['qty'], item['ref']
        total_qty += qty
        fl = FL_NC if nc else (FL_ALT if '(或' in pn else None)
        for col, v in enumerate([pn, desc, qty, ref], 1):
            cell = ws.cell(r, col, v)
            cell.font = F_BODY; cell.border = BRD
            cell.alignment = A_CTR if col == 3 else A_LFT
            if fl: cell.fill = fl

    tr = len(items) + 2
    for col, v in enumerate(['Total', None, total_qty, None], 1):
        cell = ws.cell(tr, col, v)
        cell.font = F_TOT; cell.fill = FL_TOT
        cell.border = BRD; cell.alignment = A_CTR

    lg = tr + 2
    ws.cell(lg, 1, '圖例說明：').font = Font(name='Arial', bold=True, size=9)
    for i, (label, fl, note) in enumerate([
        ('黃底', FL_NC,  '原本為 NC 的零件'),
        ('綠底', FL_ALT, '含替代料 (P/N 欄位有「或 XXXXX」)'),
    ]):
        c = ws.cell(lg+1+i, 1, label)
        c.fill = fl; c.font = Font(name='Arial', size=9)
        ws.cell(lg+1+i, 2, note).font = Font(name='Arial', size=9)

    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 7
    ws.column_dimensions['D'].width = 70
    ws.freeze_panes = 'A2'

    wb.save(out_path)
    print(f"✓ 輸出: {out_path}  |  {len(items)} 項  |  Total Qty = {total_qty}")


# ════════════════════════════════════════════════════════════════════════════════
# 7. GUI
# ════════════════════════════════════════════════════════════════════════════════

class _StdoutToText:
    """Redirect print() into a tk.Text widget (thread-safe via after())."""
    def __init__(self, widget, root):
        self.widget = widget
        self.root   = root

    def write(self, s: str):
        def _insert():
            self.widget.configure(state='normal')
            self.widget.insert('end', s)
            self.widget.see('end')
            self.widget.configure(state='disabled')
        self.root.after(0, _insert)

    def flush(self): pass


class BomMatcherApp(tk.Tk if _HAS_TK else object):

    PAD = 8

    def __init__(self):
        super().__init__()
        self.title(f'BOM Matcher  v{VERSION}')
        self.resizable(True, True)
        self.minsize(680, 480)
        self._build_ui()
        self.after(100, self._auto_detect)   # run after mainloop starts

    # ── UI construction ───────────────────────────────────────────────────────

    def _build_ui(self):
        P = self.PAD
        self.configure(padx=P, pady=P)

        # ── 目前目錄 ──
        cwd_frame = tk.LabelFrame(self, text='目前工作目錄', padx=P, pady=4)
        cwd_frame.pack(fill='x', pady=(0, P))
        self._lbl_cwd = tk.Label(cwd_frame, text=str(Path.cwd()),
                                 anchor='w', fg='#333')
        self._lbl_cwd.pack(fill='x')

        # ── 三個路徑欄位 ──
        file_frame = tk.LabelFrame(self, text='檔案設定', padx=P, pady=P)
        file_frame.pack(fill='x', pady=(0, P))
        file_frame.columnconfigure(1, weight=1)

        self.var_bom = tk.StringVar()
        self.var_db  = tk.StringVar()
        self.var_out = tk.StringVar()

        rows = [
            ('Text BOM (.txt)',   self.var_bom, self._browse_bom),
            ('Excel DB (.xlsx)',  self.var_db,  self._browse_db),
            ('輸出 BOM (.xlsx)', self.var_out, self._browse_out),
        ]
        for r, (label, var, cmd) in enumerate(rows):
            tk.Label(file_frame, text=label, anchor='e', width=16)\
                .grid(row=r, column=0, sticky='e', pady=3, padx=(0, 6))
            tk.Entry(file_frame, textvariable=var, width=60)\
                .grid(row=r, column=1, sticky='ew', pady=3)
            tk.Button(file_frame, text='瀏覽…', width=7, command=cmd)\
                .grid(row=r, column=2, padx=(6, 0), pady=3)

        # ── 執行按鈕 ──
        btn_frame = tk.Frame(self)
        btn_frame.pack(fill='x', pady=(0, P))
        self.btn_run = tk.Button(btn_frame, text='▶  執行', width=16,
                                 bg='#1F4E79', fg='white',
                                 font=('Arial', 11, 'bold'),
                                 relief='flat', cursor='hand2',
                                 command=self._run)
        self.btn_run.pack(side='right')

        # ── 日誌 ──
        log_frame = tk.LabelFrame(self, text='執行日誌', padx=P, pady=P)
        log_frame.pack(fill='both', expand=True)
        self.log_text = scrolledtext.ScrolledText(
            log_frame, state='disabled', height=12,
            font=('Consolas', 9), bg='#1e1e1e', fg='#d4d4d4',
            insertbackground='white', wrap='word',
        )
        self.log_text.pack(fill='both', expand=True)

    # ── Auto-detect ───────────────────────────────────────────────────────────

    def _auto_detect(self):
        cwd  = Path.cwd()
        txts  = sorted(cwd.glob('*.txt'))
        xlsxs = sorted(cwd.glob('*.xlsx'))

        self._log(f'[掃描] 目錄：{cwd}\n')

        # ── TXT ──
        if len(txts) == 1:
            self._set_bom(str(txts[0]))
            self._log(f'[自動偵測] Text BOM → {txts[0].name}')
        elif len(txts) > 1:
            self._log(f'[提示] 找到 {len(txts)} 個 TXT 檔，請手動選擇：')
            for f in txts:
                self._log(f'          {f.name}')
        else:
            self._log('[提示] 未找到 TXT BOM，請手動選擇')

        # ── XLSX ──
        if len(xlsxs) == 1:
            self._set_db(str(xlsxs[0]))
            self._log(f'[自動偵測] Excel DB  → {xlsxs[0].name}')
        elif len(xlsxs) > 1:
            self._log(f'[提示] 找到 {len(xlsxs)} 個 XLSX 檔，請手動選擇：')
            for f in xlsxs:
                self._log(f'          {f.name}')
        else:
            self._log('[提示] 未找到 Excel DB，請手動選擇')

        self._log('')

    # ── File helpers ──────────────────────────────────────────────────────────

    def _set_bom(self, path: str):
        self.var_bom.set(path)
        self._update_out_default(path)

    def _set_db(self, path: str):
        self.var_db.set(path)

    def _update_out_default(self, bom_path: str):
        if bom_path and not self.var_out.get():
            p   = Path(bom_path)
            out = p.parent / f'BOM_{p.stem}_{date.today():%Y%m%d}.xlsx'
            self.var_out.set(str(out))

    def _browse_bom(self):
        f = filedialog.askopenfilename(
            title='選擇 Text BOM',
            initialdir=Path.cwd(),
            filetypes=[('Text BOM', '*.txt'), ('All files', '*.*')],
        )
        if f: self._set_bom(f)

    def _browse_db(self):
        f = filedialog.askopenfilename(
            title='選擇 Excel 零件資料庫',
            initialdir=Path.cwd(),
            filetypes=[('Excel', '*.xlsx'), ('All files', '*.*')],
        )
        if f: self._set_db(f)

    def _browse_out(self):
        f = filedialog.asksaveasfilename(
            title='輸出 BOM 路徑',
            initialdir=Path.cwd(),
            defaultextension='.xlsx',
            filetypes=[('Excel', '*.xlsx')],
        )
        if f: self.var_out.set(f)

    # ── Run ───────────────────────────────────────────────────────────────────

    def _run(self):
        bom = self.var_bom.get().strip()
        db  = self.var_db.get().strip()
        out = self.var_out.get().strip()

        if not bom:
            messagebox.showerror('錯誤', '請選擇 Text BOM 檔案'); return
        if not db:
            messagebox.showerror('錯誤', '請選擇 Excel 零件資料庫'); return
        if not out:
            messagebox.showerror('錯誤', '請指定輸出路徑'); return

        self.btn_run.configure(state='disabled')
        self._log('=' * 55)
        threading.Thread(
            target=self._run_task, args=(bom, db, out), daemon=True
        ).start()

    def _run_task(self, bom: str, db_path: str, out: str):
        old_stdout = sys.stdout
        sys.stdout = _StdoutToText(self.log_text, self)
        try:
            print(f'BOM Matcher v{VERSION}')
            print(f'  Text BOM : {bom}')
            print(f'  Excel DB : {db_path}')
            print(f'  Output   : {out}')
            print()

            items = parse_text_bom(bom)
            print(f'✓ 解析 Text BOM: {len(items)} 項')

            db_df = read_db(db_path)
            print(f'✓ 讀取零件資料庫: {len(db_df)} 筆')

            write_excel(items, db_df, out)
            self.after(0, lambda: messagebox.showinfo(
                '完成', f'BOM 已產出：\n{out}'
            ))
        except Exception as e:
            print(f'\n✗ 錯誤：{e}')
            self.after(0, lambda: messagebox.showerror('執行錯誤', str(e)))
        finally:
            sys.stdout = old_stdout
            self.after(0, lambda: self.btn_run.configure(state='normal'))
            self.after(0, lambda: self._log(''))

    # ── Log helper ────────────────────────────────────────────────────────────

    def _log(self, msg: str):
        self.log_text.configure(state='normal')
        self.log_text.insert('end', msg + '\n')
        self.log_text.see('end')
        self.log_text.configure(state='disabled')


# ════════════════════════════════════════════════════════════════════════════════
# 8. Entry point
# ════════════════════════════════════════════════════════════════════════════════

def _cli():
    ap = argparse.ArgumentParser(
        description=f'BOM Matcher v{VERSION} — Text BOM + Excel DB → BOM Excel'
    )
    ap.add_argument('text_bom')
    ap.add_argument('excel_db')
    ap.add_argument('-o', '--output')
    ap.add_argument('-s', '--sheet')
    args = ap.parse_args()

    out = args.output or str(
        Path(args.text_bom).parent /
        f'BOM_{Path(args.text_bom).stem}_{date.today():%Y%m%d}.xlsx'
    )
    print(f'BOM Matcher v{VERSION}')
    print(f'  Text BOM : {args.text_bom}')
    print(f'  Excel DB : {args.excel_db}')
    print(f'  Output   : {out}\n')

    items = parse_text_bom(args.text_bom)
    print(f'✓ 解析 Text BOM: {len(items)} 項')
    db = read_db(args.excel_db, args.sheet)
    print(f'✓ 讀取零件資料庫: {len(db)} 筆')
    write_excel(items, db, out)


def main():
    # CLI mode when arguments are provided; GUI otherwise (e.g. double-click)
    if len(sys.argv) > 1:
        _cli()
    else:
        app = BomMatcherApp()
        app.mainloop()


if __name__ == '__main__':
    main()

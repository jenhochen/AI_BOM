#!/usr/bin/env python3
"""
bom_matcher.py  v1.0
Text BOM (Altium) + Excel Parts DB  →  Formatted BOM Excel

Usage
-----
    python bom_matcher.py <text_bom.txt> <parts_db.xlsx> [-o out.xlsx] [-s SheetName]

Output columns
--------------
    P/N  |  Description  |  Qty  |  Location-Top Layer

Matching rules
--------------
    C footprint  : capacitor  — value + package, boundary-safe match
    R footprint  : resistor   — normalised value (300R→300Ω / 10K→10KΩ / 1M→1MΩ)
                                + package, boundary-safe (prevents 0Ω⊂300Ω, 2KΩ⊂402KΩ)
    L footprint  : bead       — ohm value + "BEAD" + package
                  inductor    — uH value  + "INDUCTOR" + package
    Others       : specific part name → boundary-safe match, then token fallback
    TVS generic  : Part="TVS" handled before direct-name match to block SOD false hits
    NC detection : any '/'- (or '\\'-) separated segment equals 'NC' (case-insensitive)
    Alternatives : multiple DB hits → primary P/N + "(或 XXXXX)" for each extra
"""

import re
import argparse
from pathlib import Path
from datetime import date

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

VERSION = "1.0"

# Known standard package codes (used for fallback description)
_KNOWN_PKGS = {
    '0201','0402','0603','0805','1206',
    'SOD','SOT-723','SOT723','DFN',
    'BGA-247','BGA247','BGA-154','BGA154',
}


# ════════════════════════════════════════════════════════════════════════════════
# 1. Parse Altium text BOM
# ════════════════════════════════════════════════════════════════════════════════

def parse_text_bom(path: str) -> list[dict]:
    """
    Parse an Altium tab-delimited text BOM.
    Multi-line reference continuations (lines starting with double-tab) are merged.
    Returns list of dicts: {item, qty, ref, part, fp}
    """
    items, cur = [], None
    with open(path, encoding='utf-8', errors='replace') as f:
        in_data = False
        for raw in f:
            line = raw.rstrip('\n')
            if 'Item' in line and 'Quantity' in line and 'Reference' in line:
                in_data = True
                continue
            if not in_data or line.strip().startswith('__'):
                continue
            if line.startswith('\t\t') and cur:          # continuation refs
                extra = line.strip().rstrip(',')
                if extra:
                    cur['ref'] = cur['ref'].rstrip(',') + ',' + extra
                continue
            cols = line.split('\t')
            if len(cols) >= 5 and cols[0].strip().isdigit():
                if cur:
                    items.append(cur)
                cur = {
                    'item': int(cols[0].strip()),
                    'qty':  int(cols[1].strip()),
                    'ref':  cols[2].strip().rstrip(','),
                    'part': cols[3].strip(),
                    'fp':   cols[4].strip(),
                }
    if cur:
        items.append(cur)
    return items


# ════════════════════════════════════════════════════════════════════════════════
# 2. Read parts database
# ════════════════════════════════════════════════════════════════════════════════

def read_db(path: str, sheet=None) -> pd.DataFrame:
    kw = {'sheet_name': sheet} if sheet else {'sheet_name': 0}
    df = pd.read_excel(path, **kw)
    df['P/N']         = df['P/N'].astype(str).str.strip()
    df['Description'] = df['Description'].astype(str).str.strip()
    return df


# ════════════════════════════════════════════════════════════════════════════════
# 3. NC helpers
# ════════════════════════════════════════════════════════════════════════════════

def is_nc(part: str) -> bool:
    return any(s.strip().upper() == 'NC' for s in re.split(r'[/\\]+', part))

def strip_nc(part: str) -> str:
    segs = [s for s in re.split(r'[/\\]+', part) if s.strip().upper() != 'NC']
    return '/'.join(segs).strip('/')


# ════════════════════════════════════════════════════════════════════════════════
# 4. Value & package helpers
# ════════════════════════════════════════════════════════════════════════════════

_RES_PATTERNS = [
    (re.compile(r'([\d.]+)\s*R\b', re.I),  r'\1Ω'),   # 300R  → 300Ω
    (re.compile(r'([\d.]+)\s*K\b', re.I),  r'\1KΩ'),  # 10K   → 10KΩ
    (re.compile(r'([\d.]+)\s*M\b', re.I),  r'\1MΩ'),  # 1M    → 1MΩ
]

def normalize(val: str) -> str:
    """Strip trailing tags (_1% etc.) then normalise R/K/M suffixes."""
    v = re.sub(r'_.*$', '', val).strip()
    for pat, rep in _RES_PATTERNS:
        v2, n = pat.subn(rep, v)
        if n:
            return v2
    return v

def pkg_of(fp: str) -> str:
    m = re.search(r'(0\d{3})', fp)
    if m:
        return m.group(1)
    fu = fp.upper()
    for tag in ('SOD','SOT-723','SOT723','DFN','BGA-247','BGA247','BGA-154','BGA154'):
        if tag in fu:
            return tag
    return ''

def pkg_ok(pkg: str, desc: str) -> bool:
    if not pkg:
        return True
    return pkg.replace('-','').upper() in desc.replace('-','').upper()

def value_in_desc(val: str, desc: str) -> bool:
    """
    Boundary-safe value search: val must NOT be preceded by [0-9A-Za-z.]
    and must NOT be followed by [0-9A-Za-z].
    This prevents:
      '0Ω'  matching inside '300Ω' or '10Ω'
      '2KΩ' matching inside '402KΩ'
      '1uF' matching inside '0.1uF'  (dot excluded by look-behind)
    """
    pattern = r'(?<![0-9A-Za-z.])' + re.escape(val) + r'(?![0-9A-Za-z])'
    return bool(re.search(pattern, desc, re.IGNORECASE))


# ════════════════════════════════════════════════════════════════════════════════
# 5. Core matching
# ════════════════════════════════════════════════════════════════════════════════

def find_matches(part: str, fp: str, db: pd.DataFrame) -> list:
    """
    Return de-duplicated list of matching DB rows.

    Dispatch by footprint first character:
        C → capacitor  (boundary-safe value + package)
        R → resistor   (boundary-safe normalised value + package)
        L → bead / inductor
        * → IC / diode / connector  (TVS special → direct name → token fallback)
    """
    val  = strip_nc(part)
    pkg  = pkg_of(fp)
    vu   = val.upper()
    nrm  = normalize(val)       # resistor-normalised form (may equal val)
    fpt  = fp[0].upper() if fp else ''
    seen, matches = set(), []

    def add(row):
        k = str(row['P/N'])
        if k not in seen:
            seen.add(k)
            matches.append(row)

    for _, row in db.iterrows():
        desc = str(row['Description'])
        du   = desc.upper()

        # ── Capacitor ───────────────────────────────────────────────────────
        if fpt == 'C':
            if not pkg_ok(pkg, desc):
                continue
            if value_in_desc(val, desc):
                add(row)

        # ── Resistor ────────────────────────────────────────────────────────
        elif fpt == 'R':
            if not pkg_ok(pkg, desc):
                continue
            # Try both raw value and normalised form (boundary-safe)
            if value_in_desc(val, desc) or value_in_desc(nrm, desc):
                add(row)

        # ── Inductor / Bead ─────────────────────────────────────────────────
        elif fpt == 'L':
            if not pkg_ok(pkg, desc):
                continue
            if 'BEAD' in vu and 'BEAD' in du:
                m = re.search(r'(\d+)\s*[Oo]hm', val)
                if m and f"{m.group(1)}Ω" in desc:
                    add(row)
            else:
                m = re.match(r'([\d.]+uH)', val, re.I)
                if m and m.group(1).upper() in du and 'INDUCTOR' in du:
                    add(row)

        # ── IC / Diode / Connector / Others ─────────────────────────────────
        else:
            # ① Generic TVS must be checked BEFORE the direct-name match,
            #   otherwise 'TVS' (as a substring) would match every H/TVS row
            #   including SOD-523 packages that should be excluded for TVS0201.
            if vu == 'TVS':
                if 'TVS' in du:
                    if 'SOD' not in du and pkg in ('0201', '0402'):
                        add(row)
                    elif 'SOD' in du and 'SOD' in fp.upper():
                        add(row)
                continue  # skip ② and ③ for this row regardless

            # ② Direct boundary-safe name match
            if value_in_desc(val, desc):
                add(row)
                continue

            # ③ Token fallback: split part on separators, test each token ≥4 chars
            #    Uses plain 'in' (not boundary-safe) so prefix tokens like
            #    'SM2268XT' still match 'SM2268XT2F' in description.
            tokens = [t for t in re.split(r'[_\-/\s]+', val) if len(t) >= 4]
            if any(t.upper() in du for t in tokens):
                add(row)

    return matches


def match_item(item: dict, db: pd.DataFrame) -> tuple:
    nc   = is_nc(item['part'])
    hits = find_matches(item['part'], item['fp'], db)

    if not hits:
        v = strip_nc(item['part'])
        p = pkg_of(item['fp'])
        # Show extracted package only when it is a recognised standard size;
        # otherwise fall back to the full footprint string for clarity.
        display = p if p in _KNOWN_PKGS else item['fp']
        d = f"{v} {display}".strip()
        return '', d + (' (原本為NC)' if nc else ''), nc

    pn = str(hits[0]['P/N'])
    for h in hits[1:]:
        pn += f" (或 {h['P/N']})"

    d = str(hits[0]['Description']).strip()
    return pn, d + (' (原本為NC)' if nc else ''), nc


# ════════════════════════════════════════════════════════════════════════════════
# 6. Excel output
# ════════════════════════════════════════════════════════════════════════════════

def _border() -> Border:
    s = Side(style='thin', color='AAAAAA')
    return Border(left=s, right=s, top=s, bottom=s)

def write_excel(items: list, db: pd.DataFrame, out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'BOM'

    F_HDR  = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    FL_HDR = PatternFill('solid', start_color='1F4E79')
    F_BODY = Font(name='Arial', size=10)
    F_TOT  = Font(name='Arial', bold=True, size=10)
    FL_NC  = PatternFill('solid', start_color='FFF2CC')
    FL_ALT = PatternFill('solid', start_color='E2EFDA')
    FL_TOT = PatternFill('solid', start_color='D9E1F2')
    A_CTR  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    A_LFT  = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    BRD    = _border()

    # Header
    for col, h in enumerate(['P/N', 'Description', 'Qty', 'Location-Top Layer'], 1):
        cell = ws.cell(1, col, h)
        cell.font = F_HDR; cell.fill = FL_HDR
        cell.alignment = A_CTR; cell.border = BRD
    ws.row_dimensions[1].height = 22

    # Data rows
    total_qty = 0
    for r, item in enumerate(items, 2):
        pn, desc, nc = match_item(item, db)
        qty, ref = item['qty'], item['ref']
        total_qty += qty
        fl = FL_NC if nc else (FL_ALT if '(或' in pn else None)
        for col, v in enumerate([pn, desc, qty, ref], 1):
            cell = ws.cell(r, col, v)
            cell.font   = F_BODY
            cell.border = BRD
            cell.alignment = A_CTR if col == 3 else A_LFT
            if fl:
                cell.fill = fl

    # Total row
    tr = len(items) + 2
    for col, v in enumerate(['Total', None, total_qty, None], 1):
        cell = ws.cell(tr, col, v)
        cell.font = F_TOT; cell.fill = FL_TOT
        cell.border = BRD; cell.alignment = A_CTR

    # Legend
    lg = tr + 2
    ws.cell(lg, 1, '圖例說明：').font = Font(name='Arial', bold=True, size=9)
    for i, (label, fl, note) in enumerate([
        ('黃底', FL_NC,  '原本為 NC 的零件'),
        ('綠底', FL_ALT, '含替代料 (P/N 欄位有「或 XXXXX」)'),
    ]):
        c = ws.cell(lg + 1 + i, 1, label)
        c.fill = fl; c.font = Font(name='Arial', size=9)
        ws.cell(lg + 1 + i, 2, note).font = Font(name='Arial', size=9)

    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 7
    ws.column_dimensions['D'].width = 70
    ws.freeze_panes = 'A2'

    wb.save(out_path)
    print(f"✓ 輸出: {out_path}  |  {len(items)} 項  |  Total Qty = {total_qty}")


# ════════════════════════════════════════════════════════════════════════════════
# 7. CLI entry point
# ════════════════════════════════════════════════════════════════════════════════

def main():
    ap = argparse.ArgumentParser(
        description=f'BOM Matcher v{VERSION} — Text BOM + Excel DB → BOM Excel',
    )
    ap.add_argument('text_bom',       help='Altium text BOM (.txt)')
    ap.add_argument('excel_db',       help='Parts database (.xlsx)')
    ap.add_argument('-o', '--output', help='Output Excel path (default: auto-generated)')
    ap.add_argument('-s', '--sheet',  help='DB sheet name (default: first sheet)')
    args = ap.parse_args()

    out = args.output or str(
        Path(args.text_bom).parent /
        f"BOM_{Path(args.text_bom).stem}_{date.today():%Y%m%d}.xlsx"
    )

    print(f"BOM Matcher v{VERSION}")
    print(f"  Text BOM : {args.text_bom}")
    print(f"  Excel DB : {args.excel_db}")
    print(f"  Output   : {out}\n")

    items = parse_text_bom(args.text_bom)
    print(f"✓ 解析 Text BOM: {len(items)} 項")
    db = read_db(args.excel_db, args.sheet)
    print(f"✓ 讀取零件資料庫: {len(db)} 筆")
    write_excel(items, db, out)


if __name__ == '__main__':
    main()

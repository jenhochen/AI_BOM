#!/usr/bin/env python3
"""
bom_core_0.07.py  v0.07
BOM Matcher 核心模組（無 GUI）

Text BOM (Altium) + Excel Parts DB  →  ADATA E-BOM formatted Excel

本檔可獨立當 CLI 使用，也供下列前端 import：
    bom_matcher_C_0.07.py   Streamlit 網頁版
    bom_matcher_D_0.07.py   pywebview 原生視窗版

Usage (CLI)
-----------
    python bom_core_0.07.py                     # 全自動：掃描程式所在目錄
    python bom_core_0.07.py -o out.xlsx
    python bom_core_0.07.py BOM.txt db.xlsx -k Knowledge.xlsx -o out.xlsx
"""

import re
import sys
import argparse
from pathlib import Path
from datetime import date

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import range_boundaries

VERSION = "0.07"

# ════════════════════════════════════════════════════════════════════════════════
# 0. 程式所在目錄（v0.07：全面取代 Path.cwd()）
# ════════════════════════════════════════════════════════════════════════════════

APP_DIR = Path(__file__).resolve().parent

# 依 F:\AI_BOM 現行資料夾結構的搜尋子目錄（找不到就略過）
SUB_TXT = ('TXT', 'txt')
SUB_DB  = ('ExcelDB', 'exceldb', 'DB', 'db')
SUB_IMG = ('screenshot', 'Screenshot', 'image', 'img')
SUB_OUT = ('output', 'Output', 'out')

# 產出檔的指紋：A1 是公司名 → 是 E-BOM 產出，不是零件資料庫
_OUTPUT_MARK = 'ADATA Technology Co., Ltd.'


_KNOWN_PKGS = {
    '0201', '0402', '0603', '0805', '1206',
    'SOD', 'SOT-723', 'SOT723', 'DFN',
    'BGA-247', 'BGA247', 'BGA-154', 'BGA154',
}
# RC class: resistors (6103*) and capacitors (6104*) only
_RC_PREFIXES = ('6103', '6104')

# Description-based RC detection patterns (for items with no P/N)
_CAP_VAL_START = re.compile(r'^\d+\.?\d*\s*(pF|nF|uF)', re.IGNORECASE)
_RES_WORD      = re.compile(r'\bRES\b',                  re.IGNORECASE)
_OHM_VAL       = re.compile(r'\d+\.?\d*\s*(Ω|ohm)',      re.IGNORECASE)


# ════════════════════════════════════════════════════════════════════════════════
# 1. Parse Altium text BOM
# ════════════════════════════════════════════════════════════════════════════════

def parse_text_bom(path: str) -> list:
    items, cur = [], None
    with open(path, encoding='utf-8', errors='replace') as f:
        in_data = False
        for raw in f:
            line = raw.rstrip('\n')
            if 'Item' in line and 'Quantity' in line and 'Reference' in line:
                in_data = True; continue
            if not in_data or line.strip().startswith('__'):
                continue
            if line.startswith('\t\t') and cur:
                extra = line.strip().rstrip(',')
                if extra:
                    cur['ref'] = cur['ref'].rstrip(',') + ',' + extra
                continue
            cols = line.split('\t')
            if len(cols) >= 5 and cols[0].strip().isdigit():
                if cur: items.append(cur)
                cur = {
                    'item': int(cols[0].strip()),
                    'qty':  int(cols[1].strip()),
                    'ref':  cols[2].strip().rstrip(','),
                    'part': cols[3].strip(),
                    'fp':   cols[4].strip(),
                }
    if cur: items.append(cur)
    return items


# ════════════════════════════════════════════════════════════════════════════════
# 2. Read parts database
# ════════════════════════════════════════════════════════════════════════════════

def read_db(path: str, sheet=None) -> pd.DataFrame:
    kw = {'sheet_name': sheet} if sheet else {'sheet_name': 0}
    df = pd.read_excel(path, **kw)
    df.columns = [str(c).strip() for c in df.columns]
    if 'P/N' not in df.columns or 'Description' not in df.columns:
        raise ValueError(
            f"『{Path(path).name}』不是零件資料庫："
            f"第 1 列需要 'P/N' 與 'Description' 欄，實際為 {list(df.columns)[:6]}")
    df['P/N']         = df['P/N'].astype(str).str.strip()
    df['Description'] = df['Description'].astype(str).str.strip()
    # 濾掉空列 / NaN 列
    df = df[(df['P/N'] != '') & (df['P/N'].str.lower() != 'nan')]
    return df.reset_index(drop=True)


def is_parts_db(path) -> bool:
    """第 1 列必須同時有 P/N 與 Description，且不是 E-BOM 產出檔。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        wb.close()
    except Exception:
        return False
    vals = [str(v).strip() for v in first if v is not None]
    if not vals:
        return False
    if vals[0] == _OUTPUT_MARK:          # 是自己產出的 E-BOM
        return False
    up = [v.upper() for v in vals]
    return 'P/N' in up and 'DESCRIPTION' in up


def is_knowledge_db(path) -> bool:
    """
    第 1 列包含『品名』與『包裝』→ Knowledge 檔。
    v0.07 修正：不可只看檔名。PREVIEW_BOM_knowledge_v0.06.xlsx 檔名含
    knowledge 但其實是 E-BOM 產出檔，會載入 57 筆垃圾並產生大量假命中。
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        wb.close()
    except Exception:
        return False
    vals = [str(v).strip() for v in first if v is not None]
    if not vals or vals[0] == _OUTPUT_MARK:      # 是自己產出的 E-BOM
        return False
    return '品名' in vals and '包裝' in vals


# ════════════════════════════════════════════════════════════════════════════════
# 3. NC helpers
# ════════════════════════════════════════════════════════════════════════════════

def is_nc(part: str) -> bool:
    return any(s.strip().upper() == 'NC' for s in re.split(r'[/\\]+', part))

def strip_nc(part: str) -> str:
    segs = [s for s in re.split(r'[/\\]+', part) if s.strip().upper() != 'NC']
    return '/'.join(segs).strip('/')


# ════════════════════════════════════════════════════════════════════════════════
# 4. Value & package helpers
# ════════════════════════════════════════════════════════════════════════════════

_RES_PATTERNS = [
    (re.compile(r'([\d.]+)\s*R\b', re.I), r'\1Ω'),
    (re.compile(r'([\d.]+)\s*K\b', re.I), r'\1KΩ'),
    (re.compile(r'([\d.]+)\s*M\b', re.I), r'\1MΩ'),
]

def normalize(val: str) -> str:
    v = re.sub(r'_.*$', '', val).strip()
    for pat, rep in _RES_PATTERNS:
        v2, n = pat.subn(rep, v)
        if n: return v2
    return v

def pkg_of(fp: str) -> str:
    m = re.search(r'(0\d{3})', fp)
    if m: return m.group(1)
    fu = fp.upper()
    for tag in ('SOD', 'SOT-723', 'SOT723', 'DFN',
                'BGA-247', 'BGA247', 'BGA-154', 'BGA154'):
        if tag in fu: return tag
    return ''

def pkg_ok(pkg: str, desc: str) -> bool:
    if not pkg: return True
    return pkg.replace('-', '').upper() in desc.replace('-', '').upper()

def value_in_desc(val: str, desc: str) -> bool:
    pattern = r'(?<![0-9A-Za-z.])' + re.escape(val) + r'(?![0-9A-Za-z])'
    return bool(re.search(pattern, desc, re.IGNORECASE))


# ════════════════════════════════════════════════════════════════════════════════
# 5. Core matching
# ════════════════════════════════════════════════════════════════════════════════

def find_matches(part: str, fp: str, db: pd.DataFrame) -> list:
    val  = strip_nc(part)
    pkg  = pkg_of(fp)
    vu   = val.upper()
    nrm  = normalize(val)
    fpt  = fp[0].upper() if fp else ''
    seen, matches = set(), []

    def add(row):
        k = str(row['P/N'])
        if k not in seen:
            seen.add(k); matches.append(row)

    for _, row in db.iterrows():
        desc = str(row['Description'])
        du   = desc.upper()

        if fpt == 'C':
            if not pkg_ok(pkg, desc): continue
            if value_in_desc(val, desc): add(row)

        elif fpt == 'R':
            if not pkg_ok(pkg, desc): continue
            if value_in_desc(val, desc) or value_in_desc(nrm, desc): add(row)

        elif fpt == 'L':
            if not pkg_ok(pkg, desc): continue
            if 'BEAD' in vu and 'BEAD' in du:
                m = re.search(r'(\d+)\s*[Oo]hm', val)
                if m and f"{m.group(1)}Ω" in desc: add(row)
            else:
                m = re.match(r'([\d.]+uH)', val, re.I)
                if m and m.group(1).upper() in du and 'INDUCTOR' in du: add(row)

        else:
            if vu == 'TVS':
                if 'TVS' in du:
                    if 'SOD' not in du and pkg in ('0201', '0402'): add(row)
                    elif 'SOD' in du and 'SOD' in fp.upper(): add(row)
                continue
            if value_in_desc(val, desc):
                add(row); continue
            tokens = [t for t in re.split(r'[_\-/\s]+', val) if len(t) >= 4]
            if any(t.upper() in du for t in tokens): add(row)

    return matches


def match_item(item: dict, db: pd.DataFrame) -> tuple:
    """
    v0.07：回傳 (candidates, nc)
      candidates = [(pn, desc), ...]  ← 每個元素代表一個 P/N（含替代料）
      無命中時 candidates = [('', '推測描述')]
    """
    nc   = is_nc(item['part'])
    hits = find_matches(item['part'], item['fp'], db)

    if not hits:
        v = strip_nc(item['part'])
        p = pkg_of(item['fp'])
        display = p if p in _KNOWN_PKGS else item['fp']
        return [('', f"{v} {display}".strip())], nc

    return [(str(h['P/N']), str(h['Description']).strip()) for h in hits], nc


def classify_rcec(pn: str, desc: str = '', fp: str = '') -> str:
    """
    RC: P/N starts with 6103 (resistor) or 6104 (capacitor).
    If no P/N, detect from description text or footprint prefix (C/R).
    All else → EC.
    """
    first = pn.split(';')[0].strip()
    if first.startswith(_RC_PREFIXES):
        return 'RC'
    if not first:
        if _CAP_VAL_START.search(desc): return 'RC'
        if _RES_WORD.search(desc):      return 'RC'
        if _OHM_VAL.search(desc):       return 'RC'
        if fp and fp[0].upper() in ('C', 'R'): return 'RC'
    return 'EC'


# ════════════════════════════════════════════════════════════════════════════════
# 6. Knowledge loading & matching
# ════════════════════════════════════════════════════════════════════════════════

def load_knowledge(path: str) -> list:
    """Load Layout_and_BOM_Knowledge file → list of {pn, 品名, 包裝, note}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    entries = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 5:
            continue
        _, pn, name, pkg, note = row[0], row[1], row[2], row[3], row[4]
        if not any([pn, name, pkg]):
            continue
        if isinstance(pn, (int, float)):
            pn = str(int(pn))
        entries.append({
            'pn':   str(pn).strip()   if pn   else '',
            '品名': str(name).strip() if name else '',
            '包裝': str(pkg).strip()  if pkg  else '',
            'note': str(note).strip() if note else '',
        })
    wb.close()
    return entries


def find_knowledge_hits(pn_str: str, desc_str: str, knowledge: list) -> str:
    """
    Return merged note (newline-separated, deduplicated) for all matching
    entries. Match triggers if any ONE of: PN / 品名 / 包裝 matches.
    """
    pn_str   = str(pn_str   or '')
    desc_str = str(desc_str or '')
    bom_pns  = {p.strip() for p in pn_str.split(';') if p.strip()}
    desc_nd  = desc_str.replace('-', '').upper()

    seen_notes, notes = set(), []
    for k in knowledge or []:
        hit = False
        if k['pn'] and k['pn'] in bom_pns:
            hit = True
        if not hit and k['品名'] and k['品名'].upper() in desc_str.upper():
            hit = True
        if not hit and k['包裝']:
            if k['包裝'].replace('-', '').upper() in desc_nd:
                hit = True
        if hit and k['note'] and k['note'] not in seen_notes:
            seen_notes.add(k['note'])
            notes.append(k['note'])
    return '\n'.join(notes)


# ════════════════════════════════════════════════════════════════════════════════
# 7. v0.07 新增：BOM 列建構（替代料拆列）
# ════════════════════════════════════════════════════════════════════════════════

def build_rows(items: list, db: pd.DataFrame, knowledge: list = None,
               split_alt: bool = True) -> list:
    """
    把解析後的 BOM items + 零件DB + Knowledge 轉成「可直接寫入 Excel 的列」。

    v0.07：一列只有一個 P/N。
      有替代料時（find_matches 回傳 2 筆以上）拆成多列：
        主列   item=序號   qty=數量  ref=位置   pn=主料
        續列   item=None   qty=None  ref=None   pn=替代料   alt=True
      Total Qty 只加總主列，不會因拆列而重複計算。

    回傳 list[dict]，欄位見 ROW_FIELDS。
    """
    rows = []
    for seq, item in enumerate(items, 1):
        cands, nc = match_item(item, db)
        if not split_alt:
            cands = [(';'.join(p for p, _ in cands if p), cands[0][1])]

        pn_all = ';'.join(p for p, _ in cands if p)
        rcec   = classify_rcec(pn_all, cands[0][1], item['fp'])
        group  = len(cands) > 1

        for i, (pn, desc) in enumerate(cands):
            d = desc + (' (原本為NC)' if (nc and i == 0) else '')
            rows.append({
                'item':  seq if i == 0 else None,
                'seq':   seq,                      # 群組編號（拆列共用）
                'rcec':  rcec,
                'pn':    pn,
                'desc':  d,
                'qty':   item['qty'] if i == 0 else None,
                'ref':   item['ref'] if i == 0 else '',
                'nc':    nc,
                'alt':   i > 0,                    # 是替代料續列
                'group': group,                    # 屬於有替代料的群組
                'fp':    item['fp'],
                'part':  item['part'],
                'kn':    find_knowledge_hits(pn, d, knowledge) if knowledge else '',
            })

    # ── PCB 列：固定最後一項 ────────────────────────────────────────────────
    rows.append({
        'item': len(items) + 1, 'seq': len(items) + 1, 'rcec': 'EC',
        'pn': '', 'desc': 'PCB', 'qty': 1, 'ref': '',
        'nc': False, 'alt': False, 'group': False, 'fp': '', 'part': 'PCB',
        'kn': find_knowledge_hits('', 'PCB', knowledge) if knowledge else '',
    })
    return rows


ROW_FIELDS = ('item', 'seq', 'rcec', 'pn', 'desc', 'qty', 'ref',
              'nc', 'alt', 'group', 'fp', 'part', 'kn')


def total_qty(rows: list) -> int:
    """只加總主列（qty 非 None），替代料續列不重複計。"""
    return sum(int(r['qty']) for r in rows if r.get('qty') not in (None, ''))


def stats(rows: list) -> dict:
    main = [r for r in rows if not r['alt']]
    return {
        'items':   len(main),
        'rows':    len(rows),
        'hit':     sum(1 for r in main if r['pn']),
        'miss':    sum(1 for r in main if not r['pn']),
        'nc':      sum(1 for r in main if r['nc']),
        'altgrp':  len({r['seq'] for r in rows if r['group']}),
        'kn':      sum(1 for r in rows if r['kn']),
        'qty':     total_qty(rows),
    }


# ════════════════════════════════════════════════════════════════════════════════
# 8. Style helpers
# ════════════════════════════════════════════════════════════════════════════════

def _sd(style='thin'):
    return Side(style=style, color='000000')

def _brd(l='thin', r='thin', t='thin', b='thin') -> Border:
    def s(x): return _sd(x) if x else Side()
    return Border(left=s(l), right=s(r), top=s(t), bottom=s(b))

def _aln(h='left', v='center', wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


B_OUTER  = _brd('medium', 'medium', 'thin', 'thin')
B_TOPCAP = _brd('medium', 'medium', 'medium', 'thin')
B_BOTCAP = _brd('medium', 'medium', 'thin', 'medium')
B_BODY_L = _brd('medium', 'medium', 'thin', 'thin')
B_FOOT   = _brd('medium', 'medium', 'thin', 'thin')


def _set_box(ws, rng: str, value=None, font=None, align=None,
             fill=None, border: Border = None):
    """
    Merge cells (if range), distribute border correctly across all edge
    cells, and write value/font/align/fill to the anchor cell.
    """
    c1, r1, c2, r2 = range_boundaries(rng)
    if (c1, r1) != (c2, r2):
        ws.merge_cells(rng)

    if border is not None:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = Border(
                    left  =border.left   if c == c1 else Side(),
                    right =border.right  if c == c2 else Side(),
                    top   =border.top    if r == r1 else Side(),
                    bottom=border.bottom if r == r2 else Side(),
                )
    if fill is not None:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = fill

    anchor = ws.cell(r1, c1)
    if value is not None: anchor.value     = value
    if font  is not None: anchor.font      = font
    if align is not None: anchor.alignment = align
    return anchor


# ════════════════════════════════════════════════════════════════════════════════
# 9. Header (rows 1-8)
# ════════════════════════════════════════════════════════════════════════════════

HEADER_FIELDS = ('pcb_no', 'serial_no', 'pcb_pn', 'issue_date',
                 'rcec_no', 'rev', 'pcb_vendor', 'pcb_spec', 'ic_type')

def empty_header() -> dict:
    d = {k: '' for k in HEADER_FIELDS}
    d['rev'] = '0'
    return d


def _build_header(ws, hi: dict = None):
    hi = {**empty_header(), **(hi or {})}
    FA = lambda sz, bold=False: Font(name='Arial',  size=sz, bold=bold)
    FK = lambda sz:             Font(name='標楷體', size=sz)
    R, L, C = _aln('right'), _aln('left'), _aln('center')

    ws.merge_cells('A1:O1')
    _set_box(ws, 'A1', 'ADATA Technology Co., Ltd.', FA(14, True), L)
    ws.row_dimensions[1].height = 17.4

    ws.merge_cells('A2:L2')
    _set_box(ws, 'A2', 'M.2 SSD E-BOM', FA(24, True), C)
    ws.row_dimensions[2].height = 30.6

    ws.row_dimensions[3].height = 19.8
    _set_box(ws, 'A3:C3', 'PCB板號：',  FA(14), R, border=B_TOPCAP)
    _set_box(ws, 'D3:J3', hi['pcb_no'] or None,    FA(14), L, border=B_TOPCAP)
    _set_box(ws, 'K3:L3', '發行序號：', FK(14), R, border=B_TOPCAP)
    _set_box(ws, 'M3:O3', hi['serial_no'] or None, FA(14), L, border=B_TOPCAP)

    std = [
        (4, 19.8, 'PCB料號：',   hi['pcb_pn'],     '發行(改版)日期：', hi['issue_date']),
        (5, 39.6, 'RC/EC NO.：', hi['rcec_no'],    '目前版次：',        hi['rev']),
        (6, 19.8, 'PCB供應商：', hi['pcb_vendor'], None,                None),
    ]
    for row, h, lbl_ac, val_dj, lbl_kl, val_mo in std:
        ws.row_dimensions[row].height = h
        _set_box(ws, f'A{row}:C{row}', lbl_ac,       FA(14), R, border=B_OUTER)
        _set_box(ws, f'D{row}:J{row}', val_dj or None, FA(14), L, border=B_OUTER)
        _set_box(ws, f'K{row}:L{row}', lbl_kl,       FK(14), R, border=B_OUTER)
        _set_box(ws, f'M{row}:O{row}', val_mo or None, FA(14), L, border=B_OUTER)

    ws.row_dimensions[7].height = 19.8
    ws.row_dimensions[8].height = 20.4
    _set_box(ws, 'A7:C7', 'PCB 規格：',        FA(14), R, border=B_OUTER)
    _set_box(ws, 'D7:J7', hi['pcb_spec'] or None, FA(14), L, border=B_OUTER)
    _set_box(ws, 'A8:C8', 'Support IC Type：', FA(14), R, border=B_BOTCAP)
    _set_box(ws, 'D8:J8', hi['ic_type'] or None,  FA(14), L, border=B_BOTCAP)

    box78 = _brd('medium', 'medium', 'thin', 'medium')
    _set_box(ws, 'K7:L8', '環保等級：', FK(14), R, border=box78)
    _set_box(ws, 'M7:O8',
             '□A級 □S級 ■H級 □G級\n□OTHER 特殊客戶要求  :  ',
             FK(14), _aln('left', 'center', True), border=box78)


# ════════════════════════════════════════════════════════════════════════════════
# 10. PCB image (rows 9-35)
# ════════════════════════════════════════════════════════════════════════════════

def _add_pcb_image(ws, img_path, log=print):
    for r in range(9, 36):
        ws.row_dimensions[r].height = 28.2
        _set_box(ws, f'A{r}', None, border=_brd('medium', None, None, None))
        _set_box(ws, f'O{r}', None, border=_brd(None, 'medium', None, None))

    if not img_path:
        return
    p = Path(img_path)
    if not p.exists():
        log(f"[警告] 找不到 PCB 圖片: {img_path}")
        return
    try:
        img = XLImage(str(p))
        orig_w, orig_h = img.width, img.height
        target_h = 700
        scale = target_h / orig_h
        img.height = target_h
        img.width  = int(orig_w * scale)
        ws.add_image(img, 'A9')
        log(f"✓ PCB 圖片已嵌入: {p.name}  ({img.width}×{img.height}px)")
    except Exception as e:
        log(f"[警告] PCB 圖片嵌入失敗: {e}")


# ════════════════════════════════════════════════════════════════════════════════
# 11. BOM table
# ════════════════════════════════════════════════════════════════════════════════

def _build_bom_header(ws, row: int):
    """Row 36 column headers. G:O = 'Knowledge Note' (single-side PCB)."""
    F  = Font(name='Arial', size=12, bold=True)
    BC = _aln('center')
    B  = _brd('medium', 'medium', 'medium', 'thin')
    for col, val in [('A', 'Item'), ('B', 'RC/EC'), ('C', 'P/N'),
                     ('D', 'Description'), ('E', 'Qty'),
                     ('F', 'Location-Top Layer')]:
        _set_box(ws, f'{col}{row}', val, F, BC, border=B)
    _set_box(ws, f'G{row}:O{row}', 'Knowledge Note', F, BC, border=B)
    ws.row_dimensions[row].height = 15.0


FL_NC  = PatternFill('solid', start_color='FFF2CC')   # yellow  = NC
FL_ALT = PatternFill('solid', start_color='E2EFDA')   # green   = alt P/N group
FL_KN  = PatternFill('solid', start_color='FFE0B2')   # orange  = knowledge hit
FL_TOT = PatternFill('solid', start_color='D9E1F2')   # blue    = total


def _write_data_row(ws, row: int, r: dict):
    F_A = Font(name='Arial', size=14)
    F_B = Font(name='Arial', size=12)
    F_N = Font(name='Arial', size=10)

    kn_note = r.get('kn') or ''
    # Fill priority: knowledge orange > NC yellow > alt-group green
    fill = (FL_KN  if kn_note
            else FL_NC  if r.get('nc')
            else FL_ALT if r.get('group')
            else None)

    qty = r.get('qty')
    data = [
        ('A', r.get('item'),           F_A, _aln('center')),
        ('B', r.get('rcec') or '',     F_B, _aln('center')),
        ('C', r.get('pn') or '',       F_B, _aln('left')),
        ('D', r.get('desc') or '',     F_B, _aln('left', wrap=True)),
        ('E', qty if qty not in (None, '') else None, F_B, _aln('center')),
        ('F', r.get('ref') or '',      F_B, _aln('left', wrap=True)),
    ]
    for col, val, fnt, aln in data:
        _set_box(ws, f'{col}{row}', val, fnt, aln, fill=fill, border=B_BODY_L)

    _set_box(ws, f'G{row}:O{row}', kn_note or None, F_N,
             _aln('left', wrap=True), fill=fill, border=B_BODY_L)

    lines = max(1, kn_note.count('\n') + 1) if kn_note else 1
    ws.row_dimensions[row].height = max(17.4, lines * 18.0)


def _write_total_row(ws, row: int, tq: int):
    F = Font(name='Arial', size=12, bold=True)
    _set_box(ws, f'A{row}:B{row}', 'Total', F, _aln('center'),
             fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'C{row}', None, F, _aln('center'), fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'D{row}', None, F, _aln('center'), fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'E{row}', tq, F, _aln('center'), fill=FL_TOT, border=B_BODY_L)
    _set_box(ws, f'F{row}:O{row}', None, F, _aln('center'),
             fill=FL_TOT, border=B_BODY_L)
    ws.row_dimensions[row].height = 17.4


def _write_footer(ws, row: int):
    F_SM  = Font(name='Arial', size=10)
    F_LBL = Font(name='Arial', size=11)

    row += 1
    _set_box(ws, f'A{row}:O{row}', None, border=B_FOOT)

    row += 1
    _set_box(ws, f'A{row}:O{row}',
             'Remark：Top Layer only BOM. NC items listed for reference. '
             '替代料已拆列，每列僅含一個 P/N；Qty 只記於主列，不重複計算。',
             F_SM, _aln('left'), border=B_FOOT)
    ws.row_dimensions[row].height = 14.0

    row += 1
    _set_box(ws, f'A{row}:O{row}', None, border=B_FOOT)

    row += 1
    _set_box(ws, f'A{row}', '圖例',
             Font(name='Arial', bold=True, size=9), border=B_FOOT)
    _set_box(ws, f'B{row}:O{row}', None, border=B_FOOT)

    legends = [
        ('黃底', FL_NC,  '原本為 NC 的零件'),
        ('綠底', FL_ALT, '含替代料（已拆列，Item/Qty/Location 只記於主列）'),
        ('橘底', FL_KN,  'Knowledge 命中（含注意事項）'),
    ]
    for i, (label, fl, note) in enumerate(legends):
        r = row + 1 + i
        _set_box(ws, f'A{r}', label,
                 Font(name='Arial', size=9), fill=fl, border=B_FOOT)
        _set_box(ws, f'B{r}', note,
                 Font(name='Arial', size=9), border=B_FOOT)
        _set_box(ws, f'C{r}:O{r}', None, border=B_FOOT)
    row += len(legends)

    row += 1
    _set_box(ws, f'A{row}:O{row}', None, border=B_FOOT)   # blank divider
    row += 1
    _set_box(ws, f'A{row}:B{row+1}', '    Approval  :  ', F_LBL,
             _aln('center'), border=B_FOOT)
    _set_box(ws, f'C{row}:D{row+1}', None, border=B_FOOT)
    _set_box(ws, f'E{row}:F{row+1}', 'Rechecked: ', F_LBL,
             _aln('center'), border=B_FOOT)
    _set_box(ws, f'G{row}:H{row+1}', 'Prepared: ', F_LBL,
             _aln('center'), border=B_FOOT)
    _set_box(ws, f'I{row}:J{row+1}', None, border=B_FOOT)
    _set_box(ws, f'K{row}:O{row+1}', None, border=B_FOOT)

    last = row + 1
    for col in range(1, 16):
        c = ws.cell(last, col)
        c.border = Border(left=c.border.left, right=c.border.right,
                          top=c.border.top, bottom=_sd('medium'))


def _set_col_widths(ws):
    for col, w in {
        'A': 5.1,  'B': 10.1, 'C': 29.1, 'D': 77.8, 'E': 9.2,
        'F': 54.3, 'G': 11.0, 'H': 16.0, 'I': 13.0, 'J': 13.0,
        'K':  9.0, 'L': 12.0, 'M': 13.3, 'N':  9.0, 'O': 13.0,
    }.items():
        ws.column_dimensions[col].width = w


# ════════════════════════════════════════════════════════════════════════════════
# 12. Main write function
# ════════════════════════════════════════════════════════════════════════════════

BOM_HDR = 36     # BOM 欄標題所在列（沿用 ADATA E-BOM 範本）


def write_excel(rows: list, out_path: str, pcb_img: str = None,
                header_info: dict = None, log=print):
    """v0.07：輸入已建構好的 rows（可先經 GUI 人工修正），輸出 Excel。"""
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'BOM'
    ws.sheet_view.showGridLines = False

    _set_col_widths(ws)
    _build_header(ws, header_info)
    _add_pcb_image(ws, pcb_img, log=log)
    _build_bom_header(ws, BOM_HDR)

    r = BOM_HDR + 1
    for row in rows:
        _write_data_row(ws, r, row)
        r += 1

    tq = total_qty(rows)
    _write_total_row(ws, r, tq)
    _write_footer(ws, r)

    ws.freeze_panes = None

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    s = stats(rows)
    log(f"✓ 輸出: {out_path}")
    log(f"  {s['items']} 項（含PCB）→ {s['rows']} 列"
        f"（替代料拆出 {s['rows'] - s['items']} 列）  |  Total Qty = {tq}")
    return s


# ════════════════════════════════════════════════════════════════════════════════
# 13. Auto-detect（v0.07：以程式所在目錄為基準，含子資料夾）
# ════════════════════════════════════════════════════════════════════════════════

def _dirs(base: Path, subs) -> list:
    """base + 存在的子資料夾。Windows 不分大小寫，需去重避免同一夾掃兩次。"""
    out, seen = [], set()
    for p in [base] + [base / s for s in subs]:
        if not p.is_dir():
            continue
        k = str(p.resolve()).lower()
        if k not in seen:
            seen.add(k); out.append(p)
    return out


def is_bom_txt(path) -> bool:
    """
    必須是 Altium 匯出的 Text BOM：表頭列以 TAB 分隔，且欄位中同時有
    Item / Quantity / Reference。

    v0.07 修正一：只用副檔名會把 changelog_*.txt 當成 BOM 檔。
    v0.07 修正二：只用「文字包含」仍會誤判 —— 本專案的 README 與 changelog
                  裡就寫著「內容須含 Bill Of Materials 或 Item/Quantity/
                  Reference」這句說明，結果說明文件自己命中了偵測規則。
                  改成要求 TAB 分隔的表頭列（與 parse_text_bom 實際依據一致）。
    """
    try:
        with open(path, encoding='utf-8', errors='replace') as f:
            head = f.read(8192)
    except Exception:
        return False
    for ln in head.splitlines():
        cols = [c.strip() for c in ln.split('\t')]
        if len(cols) >= 3 and 'Item' in cols and 'Quantity' in cols \
                and 'Reference' in cols:
            return True
    return False


# Altium 匯出的副檔名可能是 .txt 或 .BOM（Windows glob 不分大小寫）
_BOM_EXTS = ('*.txt', '*.BOM')


def find_bom_txts(base: Path = APP_DIR) -> list:
    found = []
    for d in _dirs(base, SUB_TXT):
        cands = []
        for ext in _BOM_EXTS:
            cands += sorted(d.glob(ext))
        found += [p for p in cands if is_bom_txt(p)]
    # 去重（Windows 不分大小寫）
    uniq, seen = [], set()
    for p in found:
        k = str(p).lower()
        if k not in seen:
            seen.add(k); uniq.append(p)
    return uniq


def find_parts_dbs(base: Path = APP_DIR) -> list:
    found = []
    for d in _dirs(base, SUB_DB):
        for p in sorted(d.glob('*.xlsx')):
            if p.name.startswith('~$'):
                continue
            if is_parts_db(p):
                found.append(p)
    return found


def find_knowledge_dbs(base: Path = APP_DIR) -> list:
    """只認表頭（項次|PN|品名|包裝|note），不認檔名 —— 見 is_knowledge_db。"""
    found = []
    for d in _dirs(base, SUB_DB):
        for p in sorted(d.glob('*.xlsx')):
            if p.name.startswith('~$'):
                continue
            if is_knowledge_db(p):
                found.append(p)
    # 檔名含 knowledge 的排後面（= discover 取 [-1] 時優先選中）
    found.sort(key=lambda p: (0 if 'knowledge' not in p.name.lower() else 1,
                              str(p).lower()))
    return found


def find_pcb_imgs(base: Path = APP_DIR) -> list:
    found = []
    for d in _dirs(base, SUB_IMG):
        for ext in ('*.jpg', '*.jpeg', '*.png'):
            found += sorted(d.glob(ext))
    # pcb.* 優先
    found.sort(key=lambda p: (0 if p.stem.lower() == 'pcb' else 1, str(p).lower()))
    uniq, seen = [], set()
    for p in found:
        k = str(p).lower()
        if k not in seen:
            seen.add(k); uniq.append(p)
    return uniq


def output_dir(base: Path = APP_DIR) -> Path:
    for s in SUB_OUT:
        p = base / s
        if p.is_dir():
            return p
    return base


def default_out_name(bom_txt: str, base: Path = APP_DIR) -> str:
    stem = Path(bom_txt).stem if bom_txt else 'BOM'
    stem = re.sub(r'\.BOM$', '', stem, flags=re.I)   # xxx.BOM.txt → xxx
    return str(output_dir(base) / f'BOM_{stem}_{date.today():%Y%m%d}.xlsx')


def discover(base: Path = APP_DIR) -> dict:
    """一次回傳所有自動偵測結果（給三個前端共用）。"""
    txts = find_bom_txts(base)
    dbs  = find_parts_dbs(base)
    kns  = find_knowledge_dbs(base)
    imgs = find_pcb_imgs(base)
    sel_txt = str(txts[-1]) if txts else ''
    return {
        'base':      str(base),
        'txts':      [str(p) for p in txts],
        'dbs':       [str(p) for p in dbs],
        'kns':       [str(p) for p in kns],
        'imgs':      [str(p) for p in imgs],
        'bom':       sel_txt,
        'db':        str(dbs[-1])  if dbs  else '',
        'knowledge': str(kns[-1])  if kns  else '',
        'pcb':       str(imgs[0])  if imgs else '',
        'out_dir':   str(output_dir(base)),
        'out':       default_out_name(sel_txt, base),
    }


# ════════════════════════════════════════════════════════════════════════════════
# 14. 一次跑完（給 GUI 呼叫）
# ════════════════════════════════════════════════════════════════════════════════

def run_match(bom_txt: str, db_path: str, kn_path: str = '',
              split_alt: bool = True, log=print) -> tuple:
    """回傳 (rows, items, db_df, knowledge)。"""
    items = parse_text_bom(bom_txt)
    log(f'✓ 解析 Text BOM: {len(items)} 項')
    db = read_db(db_path)
    log(f'✓ 讀取零件資料庫: {len(db)} 筆')
    kn = None
    if kn_path and Path(kn_path).exists():
        kn = load_knowledge(kn_path)
        log(f'✓ 載入 Knowledge: {len(kn)} 筆')
    rows = build_rows(items, db, kn, split_alt=split_alt)
    s = stats(rows)
    log(f'✓ 比對完成: 命中 {s["hit"]} / 未命中 {s["miss"]} / '
        f'替代料群組 {s["altgrp"]} / Knowledge 命中 {s["kn"]} 列')
    return rows, items, db, kn


# ════════════════════════════════════════════════════════════════════════════════
# 15. CLI
# ════════════════════════════════════════════════════════════════════════════════

def _cli(argv=None):
    ap = argparse.ArgumentParser(
        description=f'BOM Matcher core v{VERSION}（工作目錄 = 程式所在目錄）')
    ap.add_argument('text_bom', nargs='?', help='未指定則自動偵測')
    ap.add_argument('excel_db', nargs='?', help='未指定則自動偵測')
    ap.add_argument('-o', '--output')
    ap.add_argument('-k', '--knowledge')
    ap.add_argument('-p', '--pcb')
    ap.add_argument('--no-split', action='store_true',
                    help='關閉替代料拆列（回到 v0.06 的 a;b 單列格式）')
    args = ap.parse_args(argv)

    d = discover()
    bom = args.text_bom or d['bom']
    db  = args.excel_db or d['db']
    kn  = args.knowledge if args.knowledge is not None else d['knowledge']
    pcb = args.pcb       if args.pcb       is not None else d['pcb']
    out = args.output or default_out_name(bom)

    print(f'BOM Matcher v{VERSION}  (core / CLI)')
    print(f'  程式目錄   : {APP_DIR}')
    print(f'  Text BOM   : {bom or "（未找到）"}')
    print(f'  Excel DB   : {db  or "（未找到）"}')
    print(f'  Knowledge  : {kn  or "（未找到）"}')
    print(f'  PCB 圖片   : {pcb or "（未找到）"}')
    print(f'  Output     : {out}')
    print(f'  替代料拆列 : {"否" if args.no_split else "是"}\n')

    if not bom or not db:
        print('✗ 缺少 Text BOM 或零件資料庫，請用參數指定。')
        return 1

    rows, *_ = run_match(bom, db, kn, split_alt=not args.no_split)
    write_excel(rows, out, pcb or None)
    return 0


if __name__ == '__main__':
    sys.exit(_cli())

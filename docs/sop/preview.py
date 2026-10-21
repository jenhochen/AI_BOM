#!/usr/bin/env python3
"""
從「產出的 .pptx」反向解析每個形狀的實際座標，產出 HTML 版 SOP。

兩個用途一份檔案：
  1. 可閱讀的 HTML 版 SOP —— 左側固定目錄，點了直接跳到該段
  2. 版面 QA —— 載入後 JS 自動掃描「超出投影片邊界」與「文字爆框」並標紅

為什麼不重寫一份版面：重寫的預覽只能證明「我以為的座標」沒爆框，
證明不了 pptx 裡真正的座標。直接讀 pptx 才是對的東西。

    python docs/sop/preview.py
    → docs/BOM_Matcher_v<VER>_操作SOP.html
"""
import base64
import html
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

HERE = Path(__file__).resolve().parent
DOCS = HERE.parent
# 版號最大的那一份（舊版沒搬走時，glob 的第一個可能是舊版）
DECK = max(DOCS.glob("BOM_Matcher_v*_操作SOP.pptx"),
           key=lambda q: [int(x) for x in
                          re.search(r"_v([\d.]+)_", q.name).group(1).split(".")])
VER = re.search(r"_v([\d.]+)_", DECK.name).group(1)
OUT = DOCS / f"BOM_Matcher_v{VER}_操作SOP.html"
TOC_JSON = HERE / "toc.json"
PX = 96.0 / 914400          # EMU -> px（1 吋 = 96px）


def color_of(run):
    try:
        c = run.font.color
        if c and c.type is not None and c.rgb is not None:
            return "#%s" % str(c.rgb)
    except Exception:
        pass
    return None


def fill_css(sh):
    """
    把 pptx 形狀的填色、外框、圓角翻成 CSS。

    v0.10：不做這件事的話，build.js 疊在截圖上的紅色編號圓點在 HTML 版
    會完全不見 —— 圓點是沒有文字的 autoshape，舊版一律輸出空白 box，
    於是只剩白色數字疊在白底上。
    """
    css = ""
    try:
        if sh.fill.type == 1:                     # MSO_FILL_TYPE.SOLID
            css += f"background:#{sh.fill.fore_color.rgb};"
    except Exception:
        pass
    try:
        if sh.line.fill.type == 1:
            w = sh.line.width.pt if sh.line.width else 1
            css += f"border:{max(w, 0.75):.2g}pt solid #{sh.line.color.rgb};"
    except Exception:
        pass
    try:
        t = str(sh.auto_shape_type)
        if "OVAL" in t:
            css += "border-radius:50%;"
        elif "ROUNDED_RECTANGLE" in t:
            css += "border-radius:6px;"
    except Exception:
        pass
    return css


def shape_html(sh):
    try:
        x, y = sh.left * PX, sh.top * PX
        w, h = sh.width * PX, sh.height * PX
    except TypeError:
        return ""
    style = f"left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"

    if sh.shape_type is not None and sh.shape_type == 13:        # PICTURE
        try:
            b64 = base64.b64encode(sh.image.blob).decode()
            return (f'<div class="sh pic" style="{style}">'
                    f'<img src="data:image/{sh.image.ext};base64,{b64}"></div>')
        except Exception:
            return f'<div class="sh pic" style="{style}"></div>'

    if sh.has_table:
        rows = "".join(
            "<tr>%s</tr>" % "".join(f"<td>{html.escape(c.text)}</td>"
                                    for c in r.cells)
            for r in sh.table.rows)
        return (f'<div class="sh tbl" style="{style}">'
                f"<table>{rows}</table></div>")

    if not sh.has_text_frame or not sh.text_frame.text.strip():
        return f'<div class="sh box" style="{style}{fill_css(sh)}"></div>'

    paras = []
    for p in sh.text_frame.paragraphs:
        runs = []
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else 12
            col = color_of(r) or "#1A1C2E"
            bold = "font-weight:700;" if r.font.bold else ""
            txt = html.escape(r.text)
            # pptx 裡的跳頁連結 -> HTML 錨點
            try:
                act = r.hyperlink._hlinkClick
            except Exception:
                act = None
            if act is not None and act.get("action", "").endswith("hlinksldjump"):
                rid = act.get(
                    "{http://schemas.openxmlformats.org/officeDocument/"
                    "2006/relationships}id")
                txt = f'<a href="#s{{{rid}}}" data-rid="{rid}">{txt}</a>'
            runs.append(f'<span style="font-size:{sz}pt;color:{col};{bold}">'
                        f"{txt}</span>")
        paras.append("<p>%s</p>" % ("".join(runs) or "&nbsp;"))
    return (f'<div class="sh txt" style="{style}{fill_css(sh)}">'
            f'{"".join(paras)}</div>')


def jump_map(prs, n):
    """slideN 的 rId -> 目標頁碼"""
    out = {}
    try:
        part = prs.slides[n - 1].part
        order = {s.part.partname: i + 1 for i, s in enumerate(prs.slides)}
        for rid, rel in part.rels.items():
            if rel.reltype.endswith("/slide"):
                out[rid] = order.get(rel.target_part.partname)
    except Exception:
        pass
    return out


def main():
    prs = Presentation(str(DECK))
    sw, sh_ = prs.slide_width * PX, prs.slide_height * PX
    toc = json.loads(TOC_JSON.read_text(encoding="utf-8")) \
        if TOC_JSON.exists() else []

    parts = []
    for n, slide in enumerate(prs.slides, 1):
        # 背景色可能定義在「版面」而不是投影片本身（封面與分隔頁就是），
        # python-pptx 的 slide.background 讀不到繼承來的那一層，
        # 會害白字畫在白底上整個看不見。所以兩層都要試。
        bg = "#FFFFFF"
        for src in (slide, slide.slide_layout, slide.slide_layout.slide_master):
            try:
                f = src.background.fill
                if f.type is not None and f.fore_color.rgb is not None:
                    bg = "#%s" % str(f.fore_color.rgb)
                    break
            except Exception:
                continue
        body = "".join(shape_html(s) for s in slide.shapes)
        for rid, tgt in jump_map(prs, n).items():
            body = body.replace("#s{%s}" % rid, "#s%s" % tgt)
        parts.append(
            f'<section class="wrap" id="s{n}">'
            f'<div class="no">第 {n} 頁</div>'
            f'<div class="slide" data-n="{n}" '
            f'style="width:{sw:.0f}px;height:{sh_:.0f}px;background:{bg}">'
            f"{body}</div></section>")

    nav = ['<a class="lv0" href="#s1">封面</a>',
           '<a class="lv0" href="#s2">目錄</a>']
    for e in toc:
        cls = "lv1" if e["lvl"] == 1 else "lv2"
        label = (f'{e["n"]}　{e["title"]}' if e["lvl"] == 1
                 else e["title"])
        nav.append(f'<a class="{cls}" href="#s{e["page"]}">'
                   f'{html.escape(label)}<i>{e["page"]}</i></a>')

    doc = f"""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>BOM Matcher v{VER} 操作 SOP</title><style>
*{{box-sizing:border-box}}
body{{margin:0;background:#eef0f4;color:#1A1C2E;display:flex;
     font-family:"Microsoft JhengHei","Segoe UI",sans-serif}}
nav{{position:sticky;top:0;align-self:flex-start;height:100vh;overflow:auto;
    width:270px;flex:none;background:#fff;border-right:1px solid #dfe2e8;
    padding:18px 0}}
nav h1{{font-size:15px;margin:0 16px 2px}}
nav .sub{{font-size:11px;color:#6B7180;margin:0 16px 14px}}
nav a{{display:block;text-decoration:none;color:#3C4252;padding:5px 16px;
      font-size:12.5px;line-height:1.35;position:relative}}
nav a i{{font-style:normal;float:right;color:#9aa0ae;font-size:11px}}
nav a:hover{{background:#f2f3f8;color:#4F46E5}}
nav a.lv0{{font-weight:700;font-size:13px}}
nav a.lv1{{font-weight:700;color:#4F46E5;margin-top:8px}}
nav a.lv2{{padding-left:30px;font-size:12px;color:#5b6070}}
nav a.on{{background:#EEF0FF;color:#4F46E5;font-weight:700}}
main{{flex:1;min-width:0;padding:22px 26px 60px}}
.hd{{display:flex;align-items:baseline;gap:12px;margin-bottom:14px}}
.hd h2{{font-size:19px;margin:0}}
.hd span{{font-size:12px;color:#6B7180}}
#rep{{background:#171b24;color:#c8cede;border-radius:8px;padding:11px 14px;
     margin-bottom:20px;font-size:12px;line-height:1.7;white-space:pre-wrap;
     font-family:Consolas,monospace}}
.wrap{{margin:0 0 26px;scroll-margin-top:16px}}
.no{{font-size:11px;color:#6B7180;margin-bottom:5px}}
.slide{{position:relative;box-shadow:0 2px 10px #0002;border:1px solid #dfe2e8}}
.sh{{position:absolute;box-sizing:border-box}}
.txt p{{margin:0;line-height:1.35}}
.txt a{{color:inherit;text-decoration:none;border-bottom:1px solid currentColor}}
.pic img{{width:100%;height:100%;object-fit:fill;display:block}}
.tbl table{{width:100%;border-collapse:collapse;font-size:10.5pt}}
.tbl td{{border:1px solid #ccc;padding:4px 6px;color:#333}}
.bad{{outline:3px solid #ff4d4f !important}}
@media print{{nav,#rep{{display:none}} main{{padding:0}}}}
</style></head><body>
<nav><h1>BOM Matcher 操作 SOP</h1><div class="sub">v{VER}</div>
{''.join(nav)}</nav>
<main>
<div class="hd"><h2>BOM Matcher v{VER} 操作 SOP</h2>
<span>共 {len(parts)} 頁　點左側目錄可直接跳段</span></div>
<div id="rep">掃描中…</div>
{''.join(parts)}
</main>
<script>
const out=[];
document.querySelectorAll('.slide').forEach(sl=>{{
  const n=sl.dataset.n, R=sl.getBoundingClientRect();
  sl.querySelectorAll('.sh').forEach(el=>{{
    const r=el.getBoundingClientRect();
    if(r.left<R.left-1||r.top<R.top-1||r.right>R.right+1||r.bottom>R.bottom+1){{
      el.classList.add('bad');
      out.push('#'+n+'  元素超出投影片邊界：'+(el.innerText||'(圖/框)').slice(0,34));
    }}
    if(el.classList.contains('txt')&&el.scrollHeight>el.clientHeight+2){{
      el.classList.add('bad');
      out.push('#'+n+'  文字爆框 '+el.scrollHeight+'>'+el.clientHeight+
               '：'+el.innerText.slice(0,34));
    }}
  }});
}});
const imgs=[...document.querySelectorAll('.pic img')];
const bad=imgs.filter(i=>!i.complete||i.naturalWidth===0).length;
const links=[...document.querySelectorAll('.txt a')];
const dead=links.filter(a=>!document.querySelector(a.getAttribute('href'))).length;
document.getElementById('rep').textContent =
  '投影片 '+document.querySelectorAll('.slide').length+' 頁　圖片 '+imgs.length+
  ' 張（載入失敗 '+bad+'）　目錄連結 '+links.length+' 個（指向不存在 '+dead+'）\\n'+
  (out.length? '發現 '+out.length+' 個問題：\\n'+out.join('\\n')
             : '✓ 沒有元素超出邊界、沒有文字爆框');
window.__report={{issues:out,images:imgs.length,imgFail:bad,
                 links:links.length,deadLinks:dead}};
const secs=[...document.querySelectorAll('.wrap')];
const navs=[...document.querySelectorAll('nav a')];
addEventListener('scroll',()=>{{
  let cur=secs[0];
  for(const s of secs) if(s.getBoundingClientRect().top<120) cur=s;
  navs.forEach(a=>a.classList.toggle('on',a.getAttribute('href')==='#'+cur.id));
}},{{passive:true}});
</script></body></html>"""
    OUT.write_text(doc, encoding="utf-8")
    print("OK", OUT, f"({OUT.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    sys.exit(main())
